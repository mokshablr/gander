#!/usr/bin/env python3
"""
Generates every fixture document the test suite reads.

Three tiers share these files: the JVM and Robolectric tests take them as test
resources, the instrumented tests as assets inside the test APK, and the Python
viewer tests straight off disk. One directory, one generator, so a format only
ever needs describing once.

Everything here is invented. Nothing imitates a real organisation, and no
identifier is registry-shaped: see tests/fixtures/README.md for why that rule
exists.

Output is byte-stable, so regenerating an unchanged fixture produces no diff.
PDFs get ReportLab's invariant flag; the OOXML formats are zips, so their
entry timestamps and core properties are normalised by hand afterwards.

Usage:  python3 tests/fixtures/make_fixtures.py
Needs:  reportlab python-docx openpyxl python-pptx pillow pillow-heif cryptography
"""

import hashlib
import heapq
import hmac
import io
import os
import random
import re
import shutil
import struct
import sys
import wave
import zipfile
import zlib
from datetime import datetime
from math import pi, sin
from pathlib import Path

OUT = Path(__file__).parent / "files"

# Every timestamp written into a fixture. Arbitrary, fixed, and in the past.
EPOCH = datetime(2026, 1, 1, 0, 0, 0)
ZIP_DATE = (2026, 1, 1, 0, 0, 0)

# The word test_pdf_search.py counts. It appears exactly three times in
# six-pages.pdf and nowhere else in it, in three different cases.
NEEDLE_LINES = [
    "This tenancy begins on the first of March.",
    "The Tenancy may be ended by either party.",
    "Nothing in this TENANCY limits the above.",
]


def written(path: Path) -> Path:
    print(f"  {path.relative_to(OUT.parent.parent)}  {path.stat().st_size:,} B")
    return path


# ---------------------------------------------------------------------------
# Zip normalisation, for the three OOXML formats
# ---------------------------------------------------------------------------

ISO_EPOCH = EPOCH.strftime("%Y-%m-%dT%H:%M:%SZ")


def normalize_zip(path: Path) -> None:
    """Rewrites a zip with fixed entry timestamps, order and compression.

    python-docx and friends stamp every entry with the time the file was
    written, so an unchanged fixture would still produce a diff on every run.

    docProps/core.xml is rewritten as well. openpyxl in particular sets
    dcterms:modified to the moment of the save, overwriting whatever the
    workbook properties said, so pinning it before the save does nothing.
    """
    with zipfile.ZipFile(path) as z:
        items = sorted((i.filename, z.read(i.filename)) for i in z.infolist())
    items = [
        (name, re.sub(
            rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
            rb"\g<1>" + ISO_EPOCH.encode() + rb"\g<2>",
            data,
        ) if name == "docProps/core.xml" else data)
        for name, data in items
    ]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            z.writestr(info, data)
    tmp.replace(path)


def fix_core_properties(doc) -> None:
    """Pins the created and modified times an OOXML package records."""
    cp = doc.core_properties
    cp.created = EPOCH
    cp.modified = EPOCH
    cp.last_modified_by = "Gander tests"
    cp.author = "Gander tests"
    cp.revision = 1


# ---------------------------------------------------------------------------
# PDFs
# ---------------------------------------------------------------------------

def pdfs() -> None:
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.pdfencrypt import StandardEncryption
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    def new(path: Path, pagesize=A4, **kw):
        # invariant drops the creation date and the document id, which are the
        # only two things that would otherwise differ between two identical runs
        return canvas.Canvas(str(path), pagesize=pagesize, invariant=1, **kw)

    # six-pages.pdf: the workhorse. Base-14 fonts only, so no embedded face,
    # and exactly three occurrences of the search needle.
    c = new(OUT / "six-pages.pdf")
    for n in range(1, 7):
        c.setFont("Helvetica-Bold", 18)
        c.drawString(72, 760, f"Alder Court, page {n}")
        c.setFont("Helvetica", 11)
        y = 720
        for line in [
            "A short agreement written only so that a test has something to read.",
            "Every name, address and figure in it is invented.",
        ]:
            c.drawString(72, y, line)
            y -= 18
        if n <= len(NEEDLE_LINES):
            c.drawString(72, y - 12, NEEDLE_LINES[n - 1])
        c.showPage()
    c.save()
    written(OUT / "six-pages.pdf")

    # forty-pages.pdf: long enough that the page band holds a fraction of it,
    # which is what the virtualisation and go-to-page tests need.
    c = new(OUT / "forty-pages.pdf")
    for n in range(1, 41):
        c.setFont("Helvetica-Bold", 16)
        c.drawString(72, 760, f"Section {n}")
        c.setFont("Helvetica", 11)
        c.drawString(72, 730, f"This is page {n} of forty.")
        c.showPage()
    c.save()
    written(OUT / "forty-pages.pdf")

    # embedded-font.pdf: carries its own face, so the text layer must name
    # that face rather than a generic. ReportLab ships Vera under a licence
    # that allows redistribution.
    vera = Path(pdfmetrics.__file__).parent.parent / "fonts" / "Vera.ttf"
    pdfmetrics.registerFont(TTFont("Vera", str(vera)))
    c = new(OUT / "embedded-font.pdf")
    c.setFont("Vera", 20)
    c.drawString(72, 700, "Embedded Vera, not a system face.")
    c.showPage()
    c.save()
    written(OUT / "embedded-font.pdf")

    # cjk.pdf: a CID font named but deliberately NOT embedded, so pdf.js can
    # only draw it by loading Adobe's UniGB CMap out of lib/cmaps. Without
    # those tables the text vanishes from the canvas and the text layer both,
    # silently, which is issue #21.
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = new(OUT / "cjk.pdf")
    c.setFont("STSong-Light", 22)
    c.drawString(72, 700, "你好世界")  # ni hao shi jie
    c.setFont("Helvetica", 12)
    c.drawString(72, 660, "The line above must render and be selectable.")
    c.showPage()
    c.save()
    written(OUT / "cjk.pdf")

    # mixed-width.pdf: one A4 page then one A3, to pin that both lay out to the
    # same CSS width. That normalisation was once thought to be issue #20; it is
    # not, because a page shown at one width is equally sharp whatever its paper
    # size. The blur was the single rasterisation, and dense-map.pdf tests it.
    c = new(OUT / "mixed-width.pdf", pagesize=A4)
    c.setFont("Helvetica", 24)
    c.drawString(72, 700, "A4 page")
    c.showPage()
    c.setPageSize(A3)
    c.setFont("Helvetica", 24)
    c.drawString(72, 1000, "A3 page")
    c.showPage()
    c.save()
    written(OUT / "mixed-width.pdf")

    # dense-map.pdf: A3 landscape carrying detail into every corner, which is
    # what a tube map or a site plan is and what issue #20 was reported against.
    # The tile tests need a page whose middle is not blank: a fixture with a
    # line of text at the top correlates to nothing once you zoom past it, and
    # a test comparing two blank regions agrees with itself perfectly.
    c = new(OUT / "dense-map.pdf", pagesize=landscape(A3))
    w, h = landscape(A3)
    c.setLineWidth(0.25)
    c.setStrokeColorRGB(0.78, 0.78, 0.84)
    for x in range(0, int(w), 20):
        c.line(x, 0, x, h)
    for y in range(0, int(h), 20):
        c.line(0, y, w, y)
    c.setStrokeColorRGB(0.1, 0.2, 0.7)
    c.setLineWidth(2)
    for i in range(9):
        c.line(60 + i * 130, 60, 60 + i * 130 + 300, h - 60)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 3.5)
    for x in range(0, int(w), 100):
        for y in range(0, int(h), 60):
            c.drawString(x + 2, y + 2, f"St {x}/{y}")
    c.showPage()
    c.save()
    written(OUT / "dense-map.pdf")

    # ragged-prose.pdf: issue #22. A text layer only covers its glyphs, so the
    # leading between lines, the white beside a short line, the gutter between
    # two columns and the page margins belong to no span at all, and a finger
    # dragging a selection through one of them lands on nothing. padRows() in
    # pdf.html pads the spans until they tile the page, and this is the page
    # shaped to make it work: ragged right ends, one line of three words, a
    # wide blank before a heading, a two-column block with a gutter down the
    # middle, and a row mixing 16pt with 8pt. Every other PDF fixture here is
    # full-width lines at one size, which padRows covers without trying.
    c = new(OUT / "ragged-prose.pdf")
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 780, "Ragged prose")
    c.setFont("Helvetica", 11)
    y = 748
    for line in [
        "The first line runs the whole width of the measure and then stops here.",
        "A shorter second line.",
        "Three words.",
        "The fourth line is long again, so the ragged edge above it has somewhere",
        "to be, and the space beside it belongs to no span until padRows runs.",
    ]:
        c.drawString(72, y, line)
        y -= 18

    # The wide blank. Displacement is a fraction of a span's top padding, so it
    # is invisible on an ordinary line gap and shows up on the line after this.
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 540, "After a wide blank")
    c.setFont("Helvetica", 11)
    y = 508
    for line in [
        "This heading sits below 130 points of nothing, which is where the",
        "vertical padding is largest and any error in it is largest too.",
    ]:
        c.drawString(72, y, line)
        y -= 18

    # Two columns, so a gap inside a row has to be covered as well as the two
    # margins. A selection dragged down the left column passes through it.
    y = 430
    for left, right in [
        ("First entry", "Fourteen"),
        ("Second entry", "Twenty one"),
        ("Third entry", "Three"),
        ("Fourth entry", "Eight"),
    ]:
        c.drawString(72, y, left)
        c.drawString(340, y, right)
        y -= 18

    # One row, two type sizes. Padding each item down by the same amount from
    # where it happens to end would leave a sliver under the smaller one.
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 330, "Large label")
    c.setFont("Helvetica", 8)
    c.drawString(300, 330, "and a caption beside it, set much smaller")

    # Sideways text, which padRows skips a span at a time. The rest of the page
    # still has to come out covered.
    c.saveState()
    c.rotate(90)
    c.setFont("Helvetica", 10)
    c.drawString(200, -560, "Printed sideways in the margin")
    c.restoreState()

    c.setFont("Helvetica", 11)
    c.drawString(72, 120, "A last line, well clear of everything above it.")
    c.showPage()
    c.save()
    written(OUT / "ragged-prose.pdf")

    # colours.pdf: everything night mode has to get right, in known values so a
    # test can assert exact pixels rather than "darker".
    #
    # Page 1 is a document: white paper, near-black text, a saturated heading and
    # a vector block, all of which turn over, plus a small image that must not.
    # Page 2 is a scan: one image covering the whole page, which must turn over
    # despite being an image, because that is what a scanned book is.
    #
    # Flat colours on purpose. A photograph would make the assertions sample
    # noise, and what is being pinned here is which pixels the filter reached.
    from PIL import Image, ImageDraw
    from reportlab.lib.utils import ImageReader

    def block(rgb, size=(64, 64)):
        return ImageReader(Image.new("RGB", size, rgb))

    W, H = 400, 600
    c = new(OUT / "colours.pdf", pagesize=(W, H))
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    # Up in the top-left corner, so the first zoom tile lands on it and the tile
    # path's own coordinate mapping is exercised rather than assumed.
    c.drawImage(block((255, 0, 255)), 20, 420, width=100, height=100)
    c.setFillColorRGB(20 / 255, 20 / 255, 20 / 255)
    c.setFont("Helvetica", 14)
    c.drawString(160, H - 60, "Body text")
    c.setFillColorRGB(0, 119 / 255, 199 / 255)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(160, H - 100, "Coloured heading")
    c.setFillColorRGB(30 / 255, 150 / 255, 60 / 255)
    c.rect(160, H - 200, 80, 60, stroke=0, fill=1)
    c.drawImage(block((200, 30, 30)), 40, 120, width=180, height=180)
    c.showPage()

    scan = Image.new("RGB", (200, 300), (255, 255, 255))
    for x in range(20, 180):
        for y in range(40, 60):
            scan.putpixel((x, y), (20, 20, 20))
    c.drawImage(ImageReader(scan), 0, 0, width=W, height=H)
    c.showPage()

    # Page 3 is the case that decides night mode's image rule, and the one two
    # simpler rules got wrong. All three of these are images as far as the PDF is
    # concerned, and they want three different things:
    #
    #   the chart   an opaque white background: leaving it alone puts a white
    #               rectangle on a dark page, which is what KOReader shipped and
    #               had reported back as their issue #4986
    #   the mono    a grey photograph, so no saturation to give it away; only the
    #     photo     absence of white paper separates it from a document
    #   the colour  the easy case, and the one a size rule got wrong by treating a
    #     photo     full-page photograph as a page
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, stroke=0, fill=1)

    chart = Image.new("RGB", (260, 170), (255, 255, 255))
    pen = ImageDraw.Draw(chart)
    pen.line([(30, 140), (250, 140)], fill=(20, 20, 20), width=2)
    pen.line([(30, 10), (30, 140)], fill=(20, 20, 20), width=2)
    pen.line([(30, 120), (90, 60), (150, 95), (210, 30)], fill=(0, 119, 199), width=3)
    c.drawImage(ImageReader(chart), 30, H - 220, width=260, height=170)

    mono = Image.new("RGB", (160, 120))
    rng = random.Random(11)
    for x in range(160):
        for y in range(120):
            v = rng.randint(20, 150)
            mono.putpixel((x, y), (v, v, v))
    c.drawImage(ImageReader(mono), 30, H - 380, width=160, height=120)

    colour = Image.new("RGB", (160, 120))
    for x in range(160):
        for y in range(120):
            colour.putpixel((x, y), (rng.randint(120, 255), rng.randint(20, 90),
                                     rng.randint(20, 90)))
    c.drawImage(ImageReader(colour), 220, H - 380, width=160, height=120)

    # A product shot: a coloured object on a white background, which is what most
    # photographs in a catalogue or a listing actually are. Mostly white, so the
    # "is it mostly paper" half of the rule says document; strongly coloured, so the
    # saturation half says picture. It is the case that stops that half being dropped.
    product = Image.new("RGB", (150, 110), (252, 252, 252))
    shot = ImageDraw.Draw(product)
    shot.ellipse([30, 20, 120, 90], fill=(214, 68, 24))
    shot.ellipse([52, 34, 78, 54], fill=(250, 186, 96))
    c.drawImage(ImageReader(product), 30, H - 520, width=150, height=110)

    # A second figure, over on the right. The one on the left cannot tell whether the
    # page sample is being mapped back to the page at all, because a wrong mapping
    # still lands somewhere near the top left and still reads as paper. This one is
    # far enough across that a wrong mapping reads nothing and leaves it white.
    right = Image.new("RGB", (200, 130), (255, 255, 255))
    pen2 = ImageDraw.Draw(right)
    pen2.rectangle([20, 20, 60, 110], fill=(20, 20, 20))
    pen2.rectangle([80, 55, 120, 110], fill=(20, 20, 20))
    pen2.rectangle([140, 35, 180, 110], fill=(20, 20, 20))
    c.drawImage(ImageReader(right), 210, H - 520, width=170, height=110)
    c.showPage()

    # Page 4: two photographs that overlap, which is how a collage, a watermarked
    # photo or a figure with an inset is put together. Both are pictures and all of
    # both must survive, the overlap included. Clipping every picture out of one
    # even-odd path gets this wrong - canvas plus two holes is three crossings, which
    # even-odd reads as inside - so the overlap comes back turned over while the two
    # pictures around it do not.
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColorRGB(20 / 255, 20 / 255, 20 / 255)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(30, H - 50, "Two photographs, overlapping")

    under = Image.new("RGB", (180, 130), (206, 44, 30))
    over = Image.new("RGB", (180, 130), (28, 82, 196))
    c.drawImage(ImageReader(under), 30, H - 260, width=180, height=130)
    c.drawImage(ImageReader(over), 130, H - 320, width=180, height=130)
    c.showPage()
    c.save()
    written(OUT / "colours.pdf")

    # encrypted.pdf: the standard security handler, which is what nearly every
    # protected PDF in circulation uses and the only kind pdf.js can unlock.
    enc = StandardEncryption("gander", canPrint=1)
    c = new(OUT / "encrypted.pdf", encrypt=enc)
    c.setFont("Helvetica", 18)
    c.drawString(72, 700, "Unlocked with the password gander.")
    c.showPage()
    c.save()
    written(OUT / "encrypted.pdf")

    # Named .pdf, is not one. pdf.js must say so rather than showing nothing.
    (OUT / "not-a-pdf.pdf").write_bytes(
        b"This file is named .pdf and is plain text. It is not a PDF at all.\n"
    )
    written(OUT / "not-a-pdf.pdf")


# ---------------------------------------------------------------------------
# OOXML
# ---------------------------------------------------------------------------

def docx() -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Field Survey, Willowmere", level=1)
    doc.add_paragraph(
        "A short report written only so that a test has something to render."
    )
    # A Wingdings bullet sitting in the private use area. docx.html rewrites
    # U+F000 to U+F0FF into real Unicode, because the font is not on the phone
    # and the glyph would otherwise come out as a blank box.
    p = doc.add_paragraph()
    run = p.add_run("")
    run.font.name = "Wingdings"
    run.font.size = Pt(12)
    p.add_run(" A bullet that arrives as a private use codepoint.")
    doc.add_paragraph("The paragraph after it, so ordering is checkable.")
    fix_core_properties(doc)
    doc.save(str(OUT / "report.docx"))
    normalize_zip(OUT / "report.docx")
    written(OUT / "report.docx")


def raised_runs() -> None:
    """
    Runs set as subscript and superscript, which docx-preview draws twice: a tab
    inside a subscript run, then tabs lined up against stops after it, then a
    footnote whose marker is raised by its run rather than by a style.
    """
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("Willowmere water survey", level=1)
    p = doc.add_paragraph("Nitrate in the north well, as NO")
    lowered = p.add_run("3\t")
    lowered.font.subscript = True
    p.add_run("at 4.2 mg per litre.")
    for row in (("Well", "Depth", "Reading"), ("North", "12 m", "4.2"), ("South", "9 m", "3.8")):
        line = doc.add_paragraph("\t".join(row))
        line.paragraph_format.tab_stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.CENTER)
        line.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
    p = doc.add_paragraph("Sampled in the spring")
    raised = p.add_run("MARKER")
    raised.font.superscript = True
    p.add_run(", after the thaw.")
    fix_core_properties(doc)
    target = OUT / "raised-runs.docx"
    doc.save(str(target))
    with_footnote(target, "MARKER", "Sampling began on the first dry day.")
    normalize_zip(target)
    written(target)


def with_footnote(path: Path, marker: str, text: str) -> None:
    """Turns the run reading [marker] into a reference to a footnote reading [text].

    python-docx writes no footnotes, so the part, its relationship and its type
    are added to the package by hand, as Word lays them out.
    """
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}
    assert "word/footnotes.xml" not in items, "the template has footnotes of its own"
    body = items["word/document.xml"].decode()
    assert body.count(f"<w:t>{marker}</w:t>") == 1
    items["word/document.xml"] = body.replace(
        f"<w:t>{marker}</w:t>", '<w:footnoteReference w:id="1"/>'
    ).encode()
    items["word/footnotes.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {w}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        '<w:continuationSeparator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:id="1"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    ).encode()
    rels = items["word/_rels/document.xml.rels"].decode()
    items["word/_rels/document.xml.rels"] = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" Target="footnotes.xml" Type="http://schemas.'
        'openxmlformats.org/officeDocument/2006/relationships/footnotes"/></Relationships>',
    ).encode()
    types = items["[Content_Types].xml"].decode()
    items["[Content_Types].xml"] = types.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    ).encode()
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)


# Stands where Word wrote w:lastRenderedPageBreak, its record of where a page began.
LRPB = "[[LRPB]]"


def rewrite_parts(path: Path, edit) -> None:
    """Runs edit(name, text) over every XML part of the package, keeping what it returns."""
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}
    for name, data in items.items():
        if name.endswith(".xml"):
            items[name] = edit(name, data.decode()).encode()
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)


def field_runs(code: str, saved: str) -> str:
    """A field as Word writes one: its code, then the value it showed when last saved."""
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> {code} </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f"<w:r><w:t>{saved}</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def word_pages() -> None:
    """
    Issue #47: pages where Word recorded them. A list item, a table and the last
    section each cross a page, an explicit break carries the record Word writes after
    it, a continuous section shares its page and the last section counts in roman from
    i. The footer's fields saved 1 and 6, the values every page of a shared footer had.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.add_heading("Willowmere survey", level=1)
    doc.add_paragraph("Page one begins here.")
    doc.add_paragraph("First numbered item.", style="List Number")
    p = doc.add_paragraph("Second numbered item, ", style="List Number")
    p.add_run(LRPB + "carried on to page two.")
    doc.add_paragraph("Third numbered item.", style="List Number")
    p = doc.add_paragraph("A footnote is cited here")
    p.add_run("MARKER")
    p.add_run(".")
    table = doc.add_table(rows=3, cols=2)
    rows = (("Well", "Reading"), (LRPB + "North", "4.2"), ("South", "3.8"))
    for row, texts in zip(table.rows, rows):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    p = doc.add_paragraph("Before the explicit break.")
    p.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph(LRPB + "After the explicit break.")
    doc.add_paragraph("Section one ends here.")
    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.add_paragraph("Section two shares page four.")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("An appendix numbered in roman.")
    p = doc.add_paragraph("It runs on ")
    p.add_run(LRPB + "to a second roman page.")
    doc.sections[0].footer.paragraphs[0].text = "Page [[PAGE]] of [[NUMPAGES]]"
    fix_core_properties(doc)
    target = OUT / "word-pages.docx"
    doc.save(str(target))
    with_footnote(target, "MARKER", "Counted on the page that cites it.")

    def edit(name, xml):
        xml = re.sub(r'<w:t( xml:space="preserve")?>' + re.escape(LRPB),
                     r'<w:lastRenderedPageBreak/><w:t\1>', xml)
        xml = re.sub(
            r"<w:r>(?:<w:rPr>.*?</w:rPr>)?<w:t>Page \[\[PAGE\]\] of \[\[NUMPAGES\]\]</w:t></w:r>",
            lambda m: '<w:r><w:t xml:space="preserve">Page </w:t></w:r>'
            + field_runs(r"PAGE   \* MERGEFORMAT", "1")
            + '<w:r><w:t xml:space="preserve"> of </w:t></w:r>'
            + field_runs("NUMPAGES", "6"),
            xml,
        )
        if name == "word/document.xml":
            # The last section, which the body's own sectPr describes, counts i, ii
            last = xml.rindex("<w:sectPr")
            cols = xml.index("<w:cols", last)
            xml = xml[:cols] + '<w:pgNumType w:fmt="lowerRoman" w:start="1"/>' + xml[cols:]
        return xml

    rewrite_parts(target, edit)
    with zipfile.ZipFile(target) as z:
        parts = "".join(z.read(n).decode() for n in z.namelist() if n.endswith(".xml"))
    assert LRPB not in parts and "[[PAGE]]" not in parts
    assert parts.count("<w:lastRenderedPageBreak/>") == 4
    normalize_zip(target)
    written(target)


def word_columns() -> None:
    """
    Issue #47: Word records the top of each column as it does each page. Three pages in
    Word: three columns and the section that runs on after them share page one, with a
    record at the top of every column and of that section, then a paragraph crosses to
    page two and the last section starts page three.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.add_heading("Willowmere columns", level=1)
    doc.add_paragraph("Page one begins here.")
    doc.add_section(WD_SECTION.CONTINUOUS)
    for text, last in (("First column.", False), ("Second column.", False), ("Third column.", True)):
        p = doc.add_paragraph(LRPB + text)
        if not last:
            p.add_run().add_break(WD_BREAK.COLUMN)
    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.add_paragraph(LRPB + "Back to one column.")
    p = doc.add_paragraph("A long paragraph Word ended page one in, ")
    p.add_run(LRPB + "carried on to page two.")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph(LRPB + "Section four begins page three.")
    fix_core_properties(doc)
    target = OUT / "word-columns.docx"
    doc.save(str(target))

    def edit(name, xml):
        xml = re.sub(r'<w:t( xml:space="preserve")?>' + re.escape(LRPB),
                     r'<w:lastRenderedPageBreak/><w:t\1>', xml)
        if name == "word/document.xml":
            # The second section, whose sectPr is the second in the body, is set in three
            cols = [m.start() for m in re.finditer(r"<w:cols\b", xml)][1]
            end = xml.index("/>", cols) + 2
            xml = xml[:cols] + '<w:cols w:num="3" w:space="720"/>' + xml[end:]
        return xml

    rewrite_parts(target, edit)
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert LRPB not in body
    assert body.count("<w:lastRenderedPageBreak/>") == 6
    assert body.count('w:num="3"') == 1 and body.count('w:type="column"') == 2
    normalize_zip(target)
    written(target)


def word_unrecorded() -> None:
    """
    Issue #47: pages Word began without a record, since it records one only in a run. Four
    pages in Word, as docProps/app.xml says: page two begins among twenty-four blank lines,
    page three where Word recorded it, and page four on a table row whose first cell is empty.
    Each stretch of the document is about a page and a half, so the turns land in the middle.
    """
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_heading("Willowmere ledger", level=1)
    for n in range(1, 9):
        doc.add_paragraph(f"Ledger line {n} of page one.")
    for _ in range(24):
        doc.add_paragraph()
    doc.add_paragraph("Page two follows the blank lines.")
    for n in range(1, 9):
        doc.add_paragraph(f"Ledger line {n} of page two.")
    doc.add_paragraph(LRPB + "Page three begins where Word recorded it.")
    table = doc.add_table(rows=19, cols=2)
    table.rows[0].cells[0].text = "Entry"
    table.rows[0].cells[1].text = "Reading"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for n in range(1, 19):
        cell = table.rows[n].cells[1]
        cell.text = f"Reading {n}"
        cell.add_paragraph(f"Checked {n}")
    doc.add_paragraph("After the table.")
    doc.sections[0].footer.paragraphs[0].text = "Page [[PAGE]] of [[NUMPAGES]]"
    fix_core_properties(doc)
    target = OUT / "word-unrecorded.docx"
    doc.save(str(target))

    def edit(name, xml):
        xml = re.sub(r'<w:t( xml:space="preserve")?>' + re.escape(LRPB),
                     r'<w:lastRenderedPageBreak/><w:t\1>', xml)
        xml = re.sub(
            r"<w:r>(?:<w:rPr>.*?</w:rPr>)?<w:t>Page \[\[PAGE\]\] of \[\[NUMPAGES\]\]</w:t></w:r>",
            lambda m: '<w:r><w:t xml:space="preserve">Page </w:t></w:r>'
            + field_runs(r"PAGE   \* MERGEFORMAT", "1")
            + '<w:r><w:t xml:space="preserve"> of </w:t></w:r>'
            + field_runs("NUMPAGES", "4"),
            xml,
        )
        return xml

    rewrite_parts(target, edit)
    with zipfile.ZipFile(target) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}
    body = items["word/document.xml"].decode()
    assert LRPB not in body and body.count("<w:lastRenderedPageBreak/>") == 1
    # Twenty-four blank lines and eighteen empty first cells
    assert body.count("<w:p/>") == 42 and body.count("<w:tblHeader/>") == 1
    app, n = re.subn(rb"<Pages>\d+</Pages>", b"<Pages>4</Pages>", items["docProps/app.xml"])
    assert n == 1
    items["docProps/app.xml"] = app
    with zipfile.ZipFile(target, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)
    normalize_zip(target)
    written(target)


def word_measured() -> None:
    """
    Issue #55: a file as WPS saves one, with no record of where its pages began, so Gander
    measures them. Every line is exactly 20pt and every paragraph of several lines is cut by
    line breaks, so where each page ends depends on no font: 648pt of body on Letter paper
    with inch margins holds 32 lines and 8pt over.

    Six pages: page one is 31 lines of filler, then a heading kept with its next paragraph,
    which has no room, so both go to page two; that page ends with three lines of a ten-line
    paragraph. Page three ends with 23 lines of filler and 48pt over, where a three-line
    paragraph would leave one line alone, so it goes whole to page four, which ends six rows
    into a table, each row 2px over its line for the cell's padding. Page five draws the
    table's header row again above row seven, and 24 of the 30 blank lines after the table;
    the other six go on to page six with the last line.
    """
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    for name in ("Normal", "Heading 1"):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(20)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True

    def filler(tag, count):
        for n in range(1, count + 1):
            doc.add_paragraph(f"Filler {tag}{n:02d}")

    def lines(label, count):
        run = doc.add_paragraph().add_run(f"{label} line 1")
        for n in range(2, count + 1):
            run.add_break()
            run.add_text(f"{label} line {n}")

    filler("A", 31)
    doc.add_heading("Kept heading", level=1)
    lines("Kept body", 5)
    filler("B", 23)
    lines("Cut", 10)
    filler("C", 23)
    lines("Widow", 3)
    filler("D", 21)
    table = doc.add_table(rows=13, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Entry", "Reading"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for n in range(1, 13):
        table.rows[n].cells[0].text, table.rows[n].cells[1].text = f"Row {n:02d}", f"Value {n:02d}"
    for _ in range(30):
        doc.add_paragraph()
    doc.add_paragraph("Last line")
    section.footer.paragraphs[0].text = "Page [[PAGE]] of [[NUMPAGES]]"
    fix_core_properties(doc)
    target = OUT / "word-measured.docx"
    doc.save(str(target))

    def edit(name, xml):
        return re.sub(
            r"<w:r>(?:<w:rPr>.*?</w:rPr>)?<w:t>Page \[\[PAGE\]\] of \[\[NUMPAGES\]\]</w:t></w:r>",
            lambda m: '<w:r><w:t xml:space="preserve">Page </w:t></w:r>'
            + field_runs("PAGE", "1")
            + '<w:r><w:t xml:space="preserve"> of </w:t></w:r>'
            + field_runs("NUMPAGES", "1"),
            xml,
        )

    rewrite_parts(target, edit)
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
        footer = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
    assert "lastRenderedPageBreak" not in body and "tblHeader" in body
    assert footer.count('fldCharType="begin"') == 2
    normalize_zip(target)
    written(target)


def word_odd_even() -> None:
    """
    A file set for different odd and even pages: four pages, each ended by a page break, with
    a header and footer for odd pages and another pair for even ones. test_docx_pages.py
    also takes the setting out, which leaves the even-page pair in a file that doesn't use it.
    """
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Header"
    section.footer.paragraphs[0].text = "Footer"
    section.even_page_header.paragraphs[0].text = "Even header"
    section.even_page_footer.paragraphs[0].text = "Even footer"
    for n in range(1, 5):
        p = doc.add_paragraph(f"Page {n}.")
        if n < 4:
            p.add_run().add_break(WD_BREAK.PAGE)
    fix_core_properties(doc)
    target = OUT / "word-odd-even.docx"
    doc.save(str(target))
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
        settings = z.read("word/settings.xml").decode()
    assert settings.count("<w:evenAndOddHeaders/>") == 1
    assert body.count('w:type="even"') == 2
    normalize_zip(target)
    written(target)


def measured_base(line_pt: int):
    """A Letter file with inch margins and every line exactly [line_pt], with no page records."""
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    fmt = doc.styles["Normal"].paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    doc.styles["Normal"].font.size = Pt(12)
    return doc


def word_measured_fit() -> None:
    """
    Measured pages with nothing to spare, as students and offices write them. First 30
    paragraphs of an essay in 12pt Times New Roman, double-spaced with an indented first line,
    whose type sits a whole line's spacing inside its lines; then 60 paragraphs at exactly
    20pt that open with a tab and wrap on their own, so their lines move once the tabs are
    lined up; then three lines with tab stops at 2in and, right-aligned, at 6in, and a contents
    line that ends on a right-aligned dotted tab at 6in, with no page number after it. Then a
    page break starts a page of 27 lines at exactly 24pt, and the next paragraph opens with a
    page break, as LibreOffice writes one, which must not leave a page of its own.
    """
    import random

    from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Pt

    doc = measured_base(24)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.first_line_indent = Inches(0.5)
    words = ("survey river bank stone wall gate culvert meadow ledger reading weir gauge level "
             "harbour lantern copper granite thistle willow ember tide marker field note").split()
    rng = random.Random(55)

    def sentence(count):
        return " ".join(rng.choice(words) for _ in range(count)) + "."

    for n in range(1, 31):
        doc.add_paragraph(f"Essay {n:02d} " + sentence(rng.randint(80, 200)))
    for n in range(1, 61):
        fmt = doc.add_paragraph(f"\tTabbed {n:02d} " + sentence(rng.randint(25, 70))).paragraph_format
        fmt.line_spacing_rule, fmt.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(20)
        fmt.first_line_indent = Inches(0)
    for n in range(1, 4):
        fmt = doc.add_paragraph(f"Stop {n}\tsecond column\tright").paragraph_format
        fmt.first_line_indent = Inches(0)
        fmt.tab_stops.add_tab_stop(Inches(2))
        fmt.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
    contents = doc.add_paragraph("Contents line ending on its tab")
    contents.paragraph_format.first_line_indent = Inches(0)
    contents.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    contents.add_run("\t")

    def full(paragraph):
        fmt = paragraph.paragraph_format
        fmt.line_spacing_rule, fmt.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(24)
        fmt.first_line_indent = Inches(0)
        return paragraph

    run = full(doc.add_paragraph()).add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_text("Full 01")
    for n in range(2, 28):
        full(doc.add_paragraph(f"Full {n:02d}"))
    run = full(doc.add_paragraph()).add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_text("After the break")
    fix_core_properties(doc)
    target = OUT / "word-measured-fit.docx"
    doc.save(str(target))
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert "lastRenderedPageBreak" not in body and body.count('w:type="page"') == 2
    normalize_zip(target)
    written(target)


def word_measured_notes() -> None:
    """
    Measured pages that hold footnotes: forty paragraphs of about three lines at exactly 20pt,
    every other one citing a note of two lines. A page leaves room at its foot for the notes
    its own references cite, and the notes are numbered through the document, as Word numbers
    them, not from 1 on each page.
    """
    doc = measured_base(20)
    words = "a reading taken at the weir after rain with the gauge and the level noted"
    for n in range(1, 41):
        p = doc.add_paragraph(f"Cited {n:02d} " + " ".join([words] * 3))
        if n % 2:
            p.add_run(f"NOTE{n:02d}")
    fix_core_properties(doc)
    target = OUT / "word-measured-notes.docx"
    doc.save(str(target))
    note = "This note on the reading runs to a second line, with the date, the gauge and who took it down"
    with_footnotes(target, [(f"NOTE{n:02d}", k, f"Note {k}. {note}.") for k, n in enumerate(range(1, 41, 2), start=1)])
    normalize_zip(target)
    written(target)


def with_footnotes(target: Path, notes) -> None:
    """
    Gives the saved file at [target] a footnotes part. [notes] holds (marker, id, text): the run
    that says only marker becomes a superscript reference to footnote id, as Word writes one, and
    that note says text.
    """
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with zipfile.ZipFile(target) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}
    body, written_notes = items["word/document.xml"].decode(), ""
    for marker, k, text in notes:
        assert body.count(f"<w:t>{marker}</w:t>") == 1
        body = body.replace(
            f"<w:t>{marker}</w:t>",
            f'<w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="{k}"/>',
        )
        written_notes += f'<w:footnote w:id="{k}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote>'
    items["word/document.xml"] = body.encode()
    items["word/footnotes.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {w}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        "<w:continuationSeparator/></w:r></w:p></w:footnote>" + written_notes + "</w:footnotes>"
    ).encode()
    items["word/_rels/document.xml.rels"] = items["word/_rels/document.xml.rels"].decode().replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes" Target="footnotes.xml" Type="http://schemas.'
        'openxmlformats.org/officeDocument/2006/relationships/footnotes"/></Relationships>',
    ).encode()
    items["[Content_Types].xml"] = items["[Content_Types].xml"].decode().replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    ).encode()
    with zipfile.ZipFile(target, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)


def word_measured_sections() -> None:
    """
    Measured pages whose room changes from page to page, every line exactly 20pt. A different
    first page has a header of one line, and every later page one of six, which runs 84pt past
    its margin. Then a section runs on with two-inch margins. Then a table with a header row has
    a row of 40 lines, taller than a page, and 40 rows after it that still turn where their
    pages fill.
    """
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.shared import Inches

    doc = measured_base(20)
    first = doc.sections[0]
    first.different_first_page_header_footer = True
    first.first_page_header.paragraphs[0].text = "First page header"
    first.header.paragraphs[0].text = "Header line 1"
    for k in range(2, 7):
        first.header.add_paragraph(f"Header line {k}")
    for n in range(1, 81):
        doc.add_paragraph(f"Opening {n:03d}")
    narrow = doc.add_section(WD_SECTION.CONTINUOUS)
    narrow.different_first_page_header_footer = False
    narrow.top_margin = narrow.bottom_margin = Inches(2)
    for n in range(1, 81):
        doc.add_paragraph(f"Narrow {n:03d}")
    table = doc.add_table(rows=42, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Entry", "Reading"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    run = table.rows[1].cells[0].paragraphs[0].add_run("Tall line 1")
    for n in range(2, 41):
        run.add_break()
        run.add_text(f"Tall line {n}")
    table.rows[1].cells[1].text = "Value 01"
    for n in range(2, 42):
        table.rows[n].cells[0].text, table.rows[n].cells[1].text = f"Row {n:02d}", f"Value {n:02d}"
    doc.add_paragraph("Last line")
    fix_core_properties(doc)
    target = OUT / "word-measured-sections.docx"
    doc.save(str(target))
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert "lastRenderedPageBreak" not in body and '<w:type w:val="continuous"/>' in body
    normalize_zip(target)
    written(target)


def word_measured_tails() -> None:
    """
    Measured pages that end in blank lines, every line exactly 20pt. Sixty lines with 6pt after
    each, then forty blank lines to a page break, more than the last page of the sixty has room
    for. A short page with two blank lines, which fit. Then a page of lines that look blank:
    a picture, an empty paragraph with a border and one with a ground, an empty paragraph and
    one holding only a bookmark, which draw nothing, underlined spaces, an empty numbered item,
    an empty item of a level with no number text, which docx-preview gives the browser's own
    marker, a tab with a dotted leader and an underlined tab to sign on, and forty blank lines
    to the end.
    """
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Pt
    from PIL import Image

    doc = measured_base(20)
    for n in range(1, 61):
        doc.add_paragraph(f"Fill {n:02d}").paragraph_format.space_after = Pt(6)
    for _ in range(40):
        doc.add_paragraph()
    doc.add_paragraph("Short page").paragraph_format.page_break_before = True
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("Signed for the client:").paragraph_format.page_break_before = True
    picture = io.BytesIO()
    Image.new("RGB", (40, 20), (90, 90, 90)).save(picture, format="PNG")
    picture.seek(0)
    doc.add_paragraph().add_run().add_picture(picture, width=Inches(1))
    doc.add_paragraph("BORDERED")
    doc.add_paragraph("SHADED")
    doc.add_paragraph()
    doc.add_paragraph("BOOKMARKED")
    doc.add_paragraph().add_run(" " * 24).underline = True
    doc.add_paragraph(style="List Number")
    doc.add_paragraph("MARKED")
    leader = doc.add_paragraph()
    leader.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
    leader.add_run("\t")
    sign = doc.add_paragraph()
    sign.paragraph_format.tab_stops.add_tab_stop(Inches(4))
    sign.add_run("\t").underline = True
    for _ in range(40):
        doc.add_paragraph()
    fix_core_properties(doc)
    target = OUT / "word-measured-tails.docx"
    doc.save(str(target))

    def edit(name, xml):
        if name == "word/numbering.xml":
            assert xml.count('<w:num w:numId="1">') == 1
            xml = xml.replace('<w:num w:numId="1">', '<w:abstractNum w:abstractNumId="90"><w:lvl w:ilvl="0">'
                              '<w:start w:val="1"/><w:numFmt w:val="decimal"/></w:lvl></w:abstractNum>'
                              '<w:num w:numId="1">')
            return xml.replace("</w:numbering>", '<w:num w:numId="90"><w:abstractNumId w:val="90"/></w:num></w:numbering>')
        if name != "word/document.xml":
            return xml
        for word, props in (("BORDERED", '<w:pBdr><w:bottom w:val="single" w:sz="12" w:space="1" w:color="000000"/></w:pBdr>'),
                            ("SHADED", '<w:shd w:val="clear" w:color="auto" w:fill="D9D9D9"/>'),
                            ("MARKED", '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="90"/></w:numPr>')):
            assert xml.count(f"<w:p><w:r><w:t>{word}</w:t></w:r></w:p>") == 1
            xml = xml.replace(f"<w:p><w:r><w:t>{word}</w:t></w:r></w:p>", f"<w:p><w:pPr>{props}</w:pPr></w:p>")
        assert xml.count("<w:p><w:r><w:t>BOOKMARKED</w:t></w:r></w:p>") == 1
        return xml.replace("<w:p><w:r><w:t>BOOKMARKED</w:t></w:r></w:p>",
                           '<w:p><w:bookmarkStart w:id="0" w:name="signature"/><w:bookmarkEnd w:id="0"/></w:p>')

    rewrite_parts(target, edit)
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert "lastRenderedPageBreak" not in body and body.count("<w:tab/>") == 2 and 'w:leader="dot"' in body
    assert '<w:numId w:val="90"/>' in body
    normalize_zip(target)
    written(target)


def word_measured_tables() -> None:
    """
    Measured tables whose header rows are drawn again on each page, every line exactly 20pt.
    The first has a header row of 40 lines, taller than a page, and 40 rows under it. After a
    page break, the second has a header row citing a footnote and 120 rows, which run over
    several pages, and then two paragraphs cite two more notes.
    """
    from docx.oxml import OxmlElement

    doc = measured_base(20)
    tall = doc.add_table(rows=41, cols=2)
    run = tall.rows[0].cells[0].paragraphs[0].add_run("Head line 01")
    for n in range(2, 41):
        run.add_break()
        run.add_text(f"Head line {n:02d}")
    tall.rows[0].cells[1].text = "Reading"
    tall.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for n in range(1, 41):
        tall.rows[n].cells[0].text, tall.rows[n].cells[1].text = f"Body {n:02d}", f"Value {n:02d}"
    doc.add_paragraph("Second table").paragraph_format.page_break_before = True
    cited = doc.add_table(rows=121, cols=2)
    cited.rows[0].cells[0].paragraphs[0].add_run("Entry")
    cited.rows[0].cells[0].paragraphs[0].add_run("NOTEHEAD")
    cited.rows[0].cells[1].text = "Reading"
    cited.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for n in range(1, 121):
        cited.rows[n].cells[0].text, cited.rows[n].cells[1].text = f"Item {n:03d}", f"Value {n:03d}"
    doc.add_paragraph("After the table").add_run("NOTEAFTER")
    doc.add_paragraph("Once more").add_run("NOTELAST")
    fix_core_properties(doc)
    target = OUT / "word-measured-tables.docx"
    doc.save(str(target))
    with_footnotes(target, [("NOTEHEAD", 1, "Readings are in metres."),
                            ("NOTEAFTER", 2, "Taken after rain."),
                            ("NOTELAST", 3, "Checked twice.")])
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert "lastRenderedPageBreak" not in body and body.count("<w:tblHeader/>") == 2
    normalize_zip(target)
    written(target)


def word_colours() -> None:
    """
    Issue #47's night mode for Word: ink, a coloured heading, a colour given only by
    the theme, a highlight, a shaded cell, a bordered paragraph, and three pictures. A
    photograph stays as it is; a white-backed figure and ink drawn on nothing read as
    paper and turn over with the page.
    """
    from docx import Document
    from docx.shared import Inches, RGBColor
    from PIL import Image, ImageDraw

    def png(im) -> io.BytesIO:
        out = io.BytesIO()
        im.save(out, format="PNG")
        out.seek(0)
        return out

    rng = random.Random(47)
    photo = Image.new("RGB", (64, 48))
    photo.putdata([
        (200 + rng.randrange(40), 30 + rng.randrange(20), 30 + rng.randrange(20)) if x < 32
        else (28 + rng.randrange(20), 82 + rng.randrange(20), 196 + rng.randrange(40))
        for y in range(48) for x in range(64)
    ])
    figure = Image.new("RGB", (120, 80), "white")
    ImageDraw.Draw(figure).line([(10, 70), (40, 30), (70, 50), (110, 10)], fill="black", width=3)
    ink = Image.new("RGBA", (120, 80), (0, 0, 0, 0))
    ImageDraw.Draw(ink).line([(10, 40), (60, 10), (110, 70)], fill=(0, 0, 0, 255), width=4)

    doc = Document()
    p = doc.add_paragraph()
    heading = p.add_run("Willowmere at night")
    heading.font.color.rgb = RGBColor(0x00, 0x77, 0xC7)
    p = doc.add_paragraph()
    p.add_run("Ink set in a near black.").font.color.rgb = RGBColor(0x14, 0x14, 0x14)
    doc.add_paragraph("Ink left to the default.")
    p = doc.add_paragraph()
    p.add_run("THEMED")
    p = doc.add_paragraph()
    p.add_run("HIGHLIT")
    p = doc.add_paragraph("A paragraph with a rule under it.")
    p.paragraph_format.space_after = 0
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "SHADED"
    doc.add_picture(png(photo), width=Inches(1.6))
    doc.add_picture(png(figure), width=Inches(2.4))
    doc.add_picture(png(ink), width=Inches(2.4))
    fix_core_properties(doc)
    target = OUT / "word-colours.docx"
    doc.save(str(target))

    def edit(name, xml):
        if name != "word/document.xml":
            return xml
        xml = xml.replace("<w:r><w:t>THEMED</w:t></w:r>",
                          '<w:r><w:rPr><w:color w:themeColor="accent1"/></w:rPr>'
                          "<w:t>Coloured by the theme alone.</w:t></w:r>")
        xml = xml.replace("<w:r><w:t>HIGHLIT</w:t></w:r>",
                          '<w:r><w:rPr><w:highlight w:val="yellow"/></w:rPr>'
                          "<w:t>Highlighted in yellow.</w:t></w:r>")
        xml = re.sub(r"<w:p><w:pPr>((?:(?!</w:pPr>).)*</w:pPr><w:r><w:t>A paragraph with a rule)",
                     r'<w:p><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="12" w:space="1"'
                     r' w:color="C81E1E"/></w:pBdr>\1', xml)
        xml = xml.replace("</w:tcPr>", '<w:shd w:val="clear" w:color="auto" w:fill="1E963C"/></w:tcPr>', 1)
        return xml

    rewrite_parts(target, edit)
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert 'w:themeColor="accent1"' in body and "w:highlight" in body
    assert 'w:color="C81E1E"' in body and 'w:fill="1E963C"' in body
    normalize_zip(target)
    written(target)


def word_lines() -> None:
    """
    Lines as Word sets them, in python-docx's template: the theme's Cambria at 11pt with Word's
    1.15 lines and 10pt after each paragraph. Then a blank line, spacing of at least 14pt and of
    at least 10pt, exactly 12pt, a run in Arial, three List Paragraph items, a style that asks
    for no space between its paragraphs, and a blank line whose mark is 20pt.
    """
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    long = " ".join(["The Willowmere ledger keeps one line for every reading taken at the north well."] * 5)
    doc = Document()
    doc.add_paragraph("Auto. " + long)
    doc.add_paragraph()
    doc.add_paragraph("After the blank line.")
    for name, size, rule in (("At least fourteen. ", 14, WD_LINE_SPACING.AT_LEAST),
                             ("At least ten. ", 10, WD_LINE_SPACING.AT_LEAST),
                             ("Exactly twelve. ", 12, WD_LINE_SPACING.EXACTLY)):
        spacing = doc.add_paragraph(name + long).paragraph_format
        spacing.line_spacing = Pt(size)
        spacing.line_spacing_rule = rule
    doc.add_paragraph().add_run("Arial. " + long).font.name = "Arial"
    for n in (1, 2, 3):
        doc.add_paragraph(f"List item {n}.", style="List Paragraph")
    doc.add_paragraph("After the list.")
    mark, size = OxmlElement("w:rPr"), OxmlElement("w:sz")
    size.set(qn("w:val"), "40")
    mark.append(size)
    doc.add_paragraph()._p.get_or_add_pPr().append(mark)
    doc.add_paragraph("After the tall blank line.")
    fix_core_properties(doc)
    target = OUT / "word-lines.docx"
    doc.save(str(target))
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
        styles = z.read("word/styles.xml").decode()
    assert body.count('w:lineRule="atLeast"') == 2 and body.count('w:lineRule="exact"') == 1
    assert body.count('<w:pPr><w:rPr><w:sz w:val="40"/></w:rPr></w:pPr>') == 1
    assert re.search(r'w:styleId="ListParagraph".*?<w:contextualSpacing/>', styles, re.S)
    normalize_zip(target)
    written(target)


def cell_directions() -> None:
    """
    One table row whose cells set each text direction docx-preview reads: lrTb, which in Word
    is plain left-to-right text and which it took for vertical, tbRl and btLr, which are
    vertical, and none at all.
    """
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("Willowmere readings, by the way each cell's text runs.")
    cells = (("Left to right", "lrTb"), ("Downwards", "tbRl"), ("Upwards", "btLr"), ("Unset", None))
    table = doc.add_table(rows=1, cols=len(cells))
    table.style = "Table Grid"
    for cell, (text, direction) in zip(table.rows[0].cells, cells):
        cell.text = text
        if direction:
            way = OxmlElement("w:textDirection")
            way.set(qn("w:val"), direction)
            cell._tc.get_or_add_tcPr().append(way)
    fix_core_properties(doc)
    target = OUT / "cell-directions.docx"
    doc.save(str(target))
    with zipfile.ZipFile(target) as z:
        body = z.read("word/document.xml").decode()
    assert all(body.count(f'<w:textDirection w:val="{d}"/>') == 1 for _, d in cells if d)
    normalize_zip(target)
    written(target)


SHEET_ROWS = [
    ("Item", "Quarter", "Amount"),
    ("Surveying", "Q3", 4200),
    ("Drainage", "Q3", 1850),
    ("Fencing", "Q3", 990),
]


def sheet_colours() -> None:
    """
    colours.xlsx: every way an .xlsx colours its cells that xlsx-styles.js reads, as openpyxl
    writes them. A second sheet is long enough to be drawn a piece at a time.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Color, Font, GradientFill, PatternFill

    def solid(**colour):
        return PatternFill(fill_type="solid", fgColor=Color(**colour))

    wb = Workbook()
    ws = wb.active
    ws.title = "Colours"
    for col, text in enumerate(("Item", "Amount", "Status", "Note"), start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.fill = solid(rgb="FF1F4E79")
        cell.font = Font(bold=True, color="FFFFFFFF")
    # Red to blue, set first so that every fill after it in the list comes after a gradient, which
    # SheetJS counts as three
    ws.cell(row=4, column=3, value="Gradient").fill = GradientFill(stop=("FFFF0000", "FF0000FF"))
    rows = [
        ("Plain", None, None),
        ("Yellow", solid(rgb="FFFFFF00"), None),
        # White, Background 1, Darker 15%: theme colour 0, which is white, shaded to D9D9D9
        ("Grey", solid(theme=0, tint=-0.1499984740745262), None),
        # Number 43 of Excel's palette, FFFF99
        ("Numbered", solid(indexed=43), None),
        ("Red text", None, Font(color="FFFF0000")),
        ("Bold", None, Font(bold=True)),
        # A quarter red over white
        ("Hatched", PatternFill(fill_type="lightGray", fgColor="FFFF0000", bgColor="FFFFFFFF"), None),
    ]
    for r, (text, fill, font) in enumerate(rows, start=2):
        cell = ws.cell(row=r, column=1, value=text)
        if fill:
            cell.fill = fill
        if font:
            cell.font = font
        ws.cell(row=r, column=2, value=r * 100)
    # Black text on red reads by day; turned over, red goes pale and the text light, so at night
    # the text must be made dark again
    ws.cell(row=3, column=3, value="Alert").fill = solid(rgb="FFFF0000")
    # Row 9 filled whole, as Excel fills one: the row's own style, and the cell written in it
    # given the same. B9 and C9 are not written, so only the row colours them.
    green = solid(rgb="FF92D050")
    ws.row_dimensions[9].fill = green
    ws.cell(row=9, column=1, value="Whole row").fill = green
    # Column D filled whole, under its heading: D2 to D8 and D10 are not written
    ws.column_dimensions["D"].fill = solid(rgb="FFFFC000")
    # Where the row and the column cross, a cell written with no style of its own
    ws.cell(row=9, column=4, value="Crossing")
    ws.cell(row=10, column=1, value="After")

    long = wb.create_sheet("Long")
    long.append(("Row", "Name"))
    for n in range(1, 1201):
        long.append((n, f"Name {n}"))
    long.cell(row=1101, column=2).fill = solid(rgb="FFFFFF00")

    wb.properties.created = EPOCH
    wb.properties.modified = EPOCH
    wb.properties.creator = "Gander tests"
    wb.save(str(OUT / "colours.xlsx"))
    normalize_zip(OUT / "colours.xlsx")
    written(OUT / "colours.xlsx")


def sheet_layout() -> None:
    """
    layout.xlsx: how a sheet is laid out and its cells are set, as Excel draws them and
    xlsx-styles.js reads them, as openpyxl writes them. The second sheet is hidden and the third is
    the one the workbook was saved showing.
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "The first sheet, not the one saved showing"
    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "Kept out of sight"
    hidden.sheet_state = "hidden"

    ws = wb.create_sheet("Layout")
    ws.sheet_properties.tabColor = "FFC00000"
    # 20 characters, 140 px; B hidden; C 6 characters, 42 px; the rest Excel's default 64 px
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["B"].hidden = True
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["E"].width = 10
    # Row 2 30 points, 40 px; row 4 hidden; row 6 set by hand to 12 points, 16 px
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[4].hidden = True
    ws.row_dimensions[6].height = 12
    ws.freeze_panes = "B3"

    ws["A1"] = "A title long enough to run on over the empty cells beside it"
    ws["A2"] = "Frozen row"
    ws["B3"] = "In a hidden column"
    ws["A3"] = "Row three"
    # A number too wide for column C in the General format, and a date there in a format of its own
    ws["C3"] = 1234567.891
    ws["C5"] = datetime(2026, 10, 21)
    ws["C5"].number_format = "dd-mm-yyyy"
    ws["A4"] = "In a hidden row"
    ws["D3"] = "Right"
    ws["D3"].alignment = Alignment(horizontal="right")
    ws["D5"] = "Middle"
    ws["D5"].alignment = Alignment(horizontal="center", vertical="center")
    ws["D7"] = "In"
    ws["D7"].alignment = Alignment(indent=2)
    ws["E3"] = "Text that wraps onto the lines below it"
    ws["E3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["E6"] = "Wrapped text in a row set short by hand"
    ws["E6"].alignment = Alignment(wrap_text=True)
    ws["F3"] = 42
    ws["G3"] = True
    ws["A8"] = "Up"
    ws["A8"].alignment = Alignment(text_rotation=90)
    ws["E8"] = "Down"
    ws["E8"].alignment = Alignment(text_rotation=180)
    ws["F8"] = "Stack"
    ws["F8"].alignment = Alignment(text_rotation=255)
    ws["A9"] = "Big serif"
    ws["A9"].font = Font(name="Times New Roman", size=20, bold=True)
    ws["A10"] = "Typewriter"
    ws["A10"].font = Font(name="Courier New", size=11)
    ws["A11"] = "Calibri"
    thin = Side(style="thin")
    ws["E10"] = "Boxed"
    ws["E10"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["E11"].border = Border(bottom=Side(style="medium", color="FF0000FF"))
    ws["E12"].border = Border(bottom=Side(style="double"))
    # A dashed red left side on F10, which the edge it shares with E10 draws, heavier than E10's thin
    ws["F10"].border = Border(left=Side(style="mediumDashed", color="FFFF0000"))
    ws["A12"] = "With a note"
    ws["A12"].comment = Comment("A note about this cell", "Asha")
    ws["A13"] = "A link"
    ws["A13"].hyperlink = "https://example.com/report"
    ws["A13"].font = Font(color="FF0563C1", underline="single")
    ws.merge_cells("E14:G15")
    ws["E14"] = "Merged"
    ws["E14"].alignment = Alignment(horizontal="center")
    ws["E16"] = -5
    ws["E16"].number_format = "#,##0;[Red]-#,##0"
    ws["F16"] = 5
    ws["F16"].number_format = "#,##0;[Red]-#,##0"
    # A filled cell hides the gridlines round it
    ws["G18"].fill = PatternFill(fill_type="solid", fgColor=Color(rgb="FFFFFF00"))
    ws["G18"] = "Filled"
    # Text set right that runs on to the left over F19, and is cut off at E19, which holds something
    ws["E19"] = "x"
    ws["G19"] = "Runs on leftward"
    ws["G19"].alignment = Alignment(horizontal="right")
    wb.active = 2

    plain = wb.create_sheet("No grid")
    plain.sheet_view.showGridLines = False
    plain["A1"] = "No gridlines here"
    plain["B2"] = "Boxed"
    plain["B2"].border = Border(bottom=thin)

    bare = wb.create_sheet("No headings")
    bare.sheet_view.showRowColHeaders = False
    bare["A1"] = "No letters or numbers"

    rtl = wb.create_sheet("Right to left")
    rtl.sheet_view.rightToLeft = True
    rtl["A1"] = "Start"
    rtl["B1"] = "Next"

    # Conditional formatting, one rule of each kind Gander draws, and one it leaves alone
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
    from openpyxl.styles.differential import DifferentialStyle
    rules = wb.create_sheet("Rules")
    for r, (status, amount, scale, bar, mark) in enumerate(
            [("Done", 5, 1, 10, "x"), ("Open", 50, 2, 20, "y"), ("Done", 120, 3, 30, "x"),
             ("Late", 80, 4, 40, "z"), ("Open", 30, 5, 50, "w"), ("Done", 95, 6, 60, "y")], start=1):
        rules.append((status, amount, scale, bar, mark, r))
    red = PatternFill(fill_type="solid", bgColor="FFFFC7CE")
    rules.conditional_formatting.add("B1:B6", CellIsRule(operator="greaterThan", formula=["100"], fill=red,
                                                         font=Font(color="FF9C0006")))
    rules.conditional_formatting.add("A1:A6", FormulaRule(formula=['$A1="Done"'],
                                                          fill=PatternFill(fill_type="solid", bgColor="FFC6EFCE")))
    rules.conditional_formatting.add("C1:C6", ColorScaleRule(start_type="min", start_color="FFF8696B",
                                                             mid_type="percentile", mid_value=50, mid_color="FFFFEB84",
                                                             end_type="max", end_color="FF63BE7B"))
    rules.conditional_formatting.add("D1:D6", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    rules.conditional_formatting.add("B1:B6", Rule(type="top10", rank=1, dxf=DifferentialStyle(font=Font(bold=True))))
    rules.conditional_formatting.add("A1:A6", Rule(type="containsText", operator="containsText", text="late",
                                                   formula=['NOT(ISERROR(SEARCH("late",A1)))'],
                                                   dxf=DifferentialStyle(font=Font(color="FFFF0000"))))
    rules.conditional_formatting.add("E1:E6", Rule(type="duplicateValues",
                                                   dxf=DifferentialStyle(fill=PatternFill(fill_type="solid", bgColor="FFFFEB9C"))))
    # A formula with a range in it, which Gander does not work out, and so draws nothing for
    rules.conditional_formatting.add("F1:F6", FormulaRule(formula=["COUNTIF($F$1:$F$6,F1)>0"],
                                                          fill=PatternFill(fill_type="solid", bgColor="FF000000")))

    # A range formatted as a table in Excel's default style, whose look the workbook names and never
    # holds; C2, in a band the style fills, has a fill of its own, which wins, as a cell's formatting does
    from openpyxl.worksheet.table import Table, TableStyleInfo
    tabled = wb.create_sheet("Table")
    tabled.append(("Name", "City", "Spend"))
    for row in [("Asha", "Pune", 120), ("Ben", "Delhi", 80), ("Chen", "Mumbai", 95), ("Dev", "Pune", 60)]:
        tabled.append(row)
    tabled["C2"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    listed = Table(displayName="Spending", ref="A1:C5")
    listed.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False,
                                           showFirstColumn=False, showLastColumn=False)
    tabled.add_table(listed)

    # A picture at B2, 80 by 40 px, and a chart at E2, which Gander marks rather than draws
    import io
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as SheetImage
    from PIL import Image
    art = wb.create_sheet("Art")
    picture = Image.new("RGB", (80, 40), (31, 78, 121))
    png = io.BytesIO()
    picture.save(png, format="PNG")
    png.seek(0)
    logo = SheetImage(png)
    logo.anchor = "B2"
    art.add_image(logo)
    for row in [("City", "Spend"), ("Pune", 120), ("Delhi", 80)]:
        art.append(row)
    chart = BarChart()
    chart.title = "Spend by city"
    chart.add_data(Reference(art, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(art, min_col=1, min_row=2, max_row=3))
    art.add_chart(chart, "E2")

    # Frozen above row 60: 59 rows, more than half a phone's screen
    tall = wb.create_sheet("Frozen tall")
    for r in range(1, 200):
        tall.cell(row=r, column=1, value=f"Row {r}")
    tall.freeze_panes = "A60"

    wb.properties.created = EPOCH
    wb.properties.modified = EPOCH
    wb.properties.creator = "Gander tests"
    wb.save(str(OUT / "layout.xlsx"))
    normalize_zip(OUT / "layout.xlsx")
    written(OUT / "layout.xlsx")


def xlsx() -> None:
    import csv

    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    first.title = "Summary"
    for row in SHEET_ROWS:
        first.append(row)
    second = wb.create_sheet("Detail")
    second.append(("Note", "Value"))
    second.append(("Second sheet marker", "detail-sheet"))
    third = wb.create_sheet("Notes")
    third.append(("Third sheet marker",))
    wb.properties.created = EPOCH
    wb.properties.modified = EPOCH
    wb.properties.creator = "Gander tests"
    wb.save(str(OUT / "budget.xlsx"))
    normalize_zip(OUT / "budget.xlsx")
    written(OUT / "budget.xlsx")

    with open(OUT / "budget.csv", "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(SHEET_ROWS)
    written(OUT / "budget.csv")

    # Issue #37. A CSV as most apps save one: UTF-8 with no byte order mark. Letters
    # of two, three and four bytes, the last a flag, which is two characters of four.
    with open(OUT / "utf8.csv", "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows([
            ("Word", "Language"), ("Флаг", "Russian"), ("Straße", "German"),
            ("東京", "Japanese"), ("\U0001F1EA\U0001F1FA", "Emoji"),
        ])
    written(OUT / "utf8.csv")

    # And one that is not UTF-8 at all, as older Excel saves a CSV on Windows, which
    # the fix for #37 must leave reading as it did.
    with open(OUT / "latin1.csv", "w", newline="", encoding="latin-1") as fh:
        csv.writer(fh).writerows([("Word", "Language"), ("Café", "French"), ("Grüße", "German")])
    written(OUT / "latin1.csv")


def pptx() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    titles = ["Willowmere Kickoff", "What we found", "What happens next"]
    for n, title in enumerate(titles, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = f"Slide {n} of three."
        body.add_paragraph().text = "Invented content, for rendering only."
    fix_core_properties(prs)
    prs.save(str(OUT / "deck.pptx"))
    normalize_zip(OUT / "deck.pptx")
    written(OUT / "deck.pptx")


def without_app_properties() -> None:
    """deck.pptx with no docProps/app.xml, which Google Slides does not write."""
    target = OUT / "deck-no-app-xml.pptx"
    with zipfile.ZipFile(OUT / "deck.pptx") as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(target, "w") as z:
        for name, data in items:
            if name == "docProps/app.xml":
                continue
            if name in ("[Content_Types].xml", "_rels/.rels"):
                data, n = re.subn(rb"<(Override|Relationship)\b[^>]*docProps/app\.xml[^>]*/>", b"", data)
                assert n == 1, f"{name} does not name docProps/app.xml once"
            z.writestr(name, data)
    normalize_zip(target)
    written(target)


# Custom geometry as Google Slides and WPS write it: each path on a grid of its own
# size, a shape's paths in one list, and a straight line as one segment.
FREEFORMS = {
    # A gate's outline and two wires, three paths on one grid
    "Gate": (1828800, 1371600, [
        (1000, 1000, "M 200 0 L 600 0 L 1000 500 L 600 1000 L 200 1000 Z"),
        (1000, 1000, "M 0 250 L 200 250"),
        (1000, 1000, "M 0 750 L 200 750"),
    ]),
    # A frame and its diagonal, the diagonal drawn on a grid half the size
    "Grid": (1828800, 914400, [
        (1000, 1000, "M 0 0 L 1000 0 L 1000 1000 L 0 1000 Z"),
        (500, 500, "M 0 0 L 500 500"),
    ]),
    # A rule under a heading: one path of one straight segment
    "Rule": (4572000, 12700, [
        (1000, 10, "M 0 0 L 1000 0"),
    ]),
}


def freeform_path(w: int, h: int, commands: str) -> str:
    out, tokens = [], commands.split()
    while tokens:
        op = tokens.pop(0)
        if op == "Z":
            out.append("<a:close/>")
            continue
        x, y = tokens.pop(0), tokens.pop(0)
        tag = {"M": "moveTo", "L": "lnTo"}[op]
        out.append(f'<a:{tag}><a:pt x="{x}" y="{y}"/></a:{tag}>')
    return f'<a:path w="{w}" h="{h}">{"".join(out)}</a:path>'


def freeforms() -> None:
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Willowmere gate"
    top = 1600200
    for n, (name, (cx, cy, paths)) in enumerate(FREEFORMS.items(), start=10):
        geometry = "".join(freeform_path(*p) for p in paths)
        slide.shapes._spTree.append(parse_xml(
            f'<p:sp {nsdecls("p", "a")}><p:nvSpPr><p:cNvPr id="{n}" name="{name}"/>'
            f'<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="457200" y="{top}"/>'
            f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:custGeom><a:rect b="b" l="l" r="r" t="t"/>'
            f"<a:pathLst>{geometry}</a:pathLst></a:custGeom><a:noFill/>"
            '<a:ln w="28575"><a:solidFill><a:srgbClr val="1F4E9A"/></a:solidFill></a:ln>'
            "</p:spPr></p:sp>"
        ))
        top += cy + 457200
    fix_core_properties(prs)
    prs.save(str(OUT / "freeforms.pptx"))
    normalize_zip(OUT / "freeforms.pptx")
    written(OUT / "freeforms.pptx")


def straight_lines() -> None:
    """
    Lines that lie flat or stand upright, so that their shapes have no height or
    no width: two of PowerPoint's straight connectors, and a rule in custom
    geometry, one segment in a shape of no height, as Google Slides writes one.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, ends, rgb in (("Level", (1, 1, 5, 1), "C02020"), ("Upright", (7, 1, 7, 5), "2060C0")):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(Inches(v) for v in ends))
        line.name = name
        line.line.color.rgb = RGBColor.from_string(rgb)
        line.line.width = Pt(3)
    slide.shapes._spTree.append(parse_xml(
        f'<p:sp {nsdecls("p", "a")}><p:nvSpPr><p:cNvPr id="20" name="Rule"/><p:cNvSpPr/>'
        '<p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="914400" y="5486400"/>'
        '<a:ext cx="3657600" cy="0"/></a:xfrm><a:custGeom><a:rect b="b" l="l" r="r" t="t"/>'
        f'<a:pathLst>{freeform_path(1000, 1000, "M 0 0 L 1000 0")}</a:pathLst></a:custGeom>'
        '<a:noFill/><a:ln w="38100"><a:solidFill><a:srgbClr val="208040"/></a:solidFill></a:ln>'
        "</p:spPr></p:sp>"
    ))
    fix_core_properties(prs)
    prs.save(str(OUT / "lines.pptx"))
    normalize_zip(OUT / "lines.pptx")
    written(OUT / "lines.pptx")


def wrapping() -> None:
    """
    A line too long for its box, which has to wrap between words, and a run of
    spaces, which has to keep its width. PPTXjs writes every space as a no-break
    space, so the long line broke wherever it ran out of room, mid-word.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, top, width, height, text in (
        ("Wrapped", 0.5, 3, 4, "Every word of this line stays whole when it wraps inside a narrow box"),
        ("Spaced", 5.5, 6, 1, "left        right"),
    ):
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(width), Inches(height))
        box.name = name
        box.text_frame.word_wrap = True
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    fix_core_properties(prs)
    prs.save(str(OUT / "wrapping.pptx"))
    normalize_zip(OUT / "wrapping.pptx")
    written(OUT / "wrapping.pptx")


def weights() -> None:
    """
    Plain text in a content placeholder, and a plain run beside a bold one in a text
    box. PPTXjs gives the paragraphs of both font-weight: 100, which the plain runs
    inherited and Android drew in Roboto Thin.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Weights"
    slide.placeholders[1].text_frame.text = "Plain body text"
    box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(8), Inches(1))
    box.name = "Mixed"
    paragraph = box.text_frame.paragraphs[0]
    for text, bold in (("plain ", False), ("bold", True)):
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(28)
        if bold:
            run.font.bold = True
    fix_core_properties(prs)
    prs.save(str(OUT / "weights.pptx"))
    normalize_zip(OUT / "weights.pptx")
    written(OUT / "weights.pptx")


def inherited_bold() -> None:
    """
    Text that is bold or italic only because the deck's design says so. The master's
    title style is bold, one layout's content is bold, another layout takes the bold back
    off its title and makes its body italic. PPTXjs made a run bold or italic only when
    the run itself said so. A run's own b="0" must still win, and a second-level paragraph
    must not take the first level's bold.
    """
    from lxml import etree
    from pptx import Presentation
    from pptx.oxml.ns import qn

    def defaults(style):
        level = style.find(qn("a:lvl1pPr"))
        if level is None:
            level = etree.SubElement(style, qn("a:lvl1pPr"))
        found = level.find(qn("a:defRPr"))
        return found if found is not None else etree.SubElement(level, qn("a:defRPr"))

    def list_style(placeholder):
        body = placeholder._element.find(qn("p:txBody"))
        found = body.find(qn("a:lstStyle"))
        if found is None:
            found = etree.Element(qn("a:lstStyle"))
            body.find(qn("a:bodyPr")).addnext(found)
        return found

    prs = Presentation()
    defaults(prs.slide_master.element.find(qn("p:txStyles")).find(qn("p:titleStyle"))).set("b", "1")
    content, section = prs.slide_layouts[1], prs.slide_layouts[2]
    defaults(list_style(content.placeholders.get(idx=1))).set("b", "1")
    defaults(list_style(section.placeholders.get(idx=0))).set("b", "0")
    defaults(list_style(section.placeholders.get(idx=1))).set("i", "1")

    slide = prs.slides.add_slide(content)
    title = slide.shapes.title.text_frame.paragraphs[0]
    for text, bold in (("Bold from the master ", None), ("but not this", False)):
        run = title.add_run()
        run.text = text
        if bold is not None:
            run.font.bold = bold
    body = slide.placeholders[1].text_frame
    body.text = "Bold from the layout"
    second = body.add_paragraph()
    second.text = "Plain at the second level"
    second.level = 1

    slide = prs.slides.add_slide(section)
    slide.shapes.title.text = "Regular by its layout"
    slide.placeholders[1].text_frame.text = "Italic from the layout"
    fix_core_properties(prs)
    prs.save(str(OUT / "inherited-bold.pptx"))
    normalize_zip(OUT / "inherited-bold.pptx")
    written(OUT / "inherited-bold.pptx")


def line_breaks() -> None:
    """
    A paragraph that two line breaks make three lines, in a box wide enough that none of
    them wraps. PPTXjs dropped the first line break of a paragraph that had more than one,
    so its first two lines ran together.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.name = "Broken"
    box.text_frame.text = "The first line\vthe second\vand the third"
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(28)
    fix_core_properties(prs)
    prs.save(str(OUT / "line-breaks.pptx"))
    normalize_zip(OUT / "line-breaks.pptx")
    written(OUT / "line-breaks.pptx")


def unwrapped() -> None:
    """
    Text boxes set not to wrap (wrap="none"), each far too narrow for its line: one
    aligned left, one centred and one aligned right, beside one that wraps as usual. A
    title that takes the setting from its layout, and a text box the layout draws behind
    the slide. PPTXjs never read the setting and wrapped every one of them at its width.
    """
    from lxml import etree
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    prs = Presentation()
    layout = prs.slide_layouts[5]
    layout.placeholders.get(idx=0)._element.find(qn("p:txBody")).find(qn("a:bodyPr")).set("wrap", "none")
    behind = etree.SubElement(layout.shapes._spTree, qn("p:sp"))
    behind.append(etree.fromstring(
        '<p:nvSpPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:cNvPr id="100" name="Behind"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
    ))
    behind.append(etree.fromstring(
        '<p:spPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:xfrm><a:off x="{Inches(0.5)}" y="{Inches(6.5)}"/><a:ext cx="{Inches(1)}" cy="{Inches(0.5)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
    ))
    behind.append(etree.fromstring(
        '<p:txBody xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:bodyPr wrap="none"/><a:lstStyle/>'
        '<a:p><a:r><a:rPr lang="en-US" sz="2000"/><a:t>Drawn behind the slide and left whole</a:t></a:r></a:p>'
        '</p:txBody>'
    ))

    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(4), Inches(0.3), Inches(2), Inches(1)
    title.name = "Inheriting"
    title.text = "A title its layout leaves whole"
    for name, left, top, align, wrap, text in (
        ("Left", 0.5, 2, PP_ALIGN.LEFT, False, "Aligned left and left whole"),
        ("Centred", 4.5, 3, PP_ALIGN.CENTER, False, "Centred and left whole"),
        ("Right", 8.5, 4, PP_ALIGN.RIGHT, False, "Aligned right and left whole"),
        ("Wrapping", 0.5, 5, PP_ALIGN.LEFT, True, "This one wraps as before"),
    ):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(1), Inches(0.5))
        box.name = name
        box.text_frame.word_wrap = wrap
        paragraph = box.text_frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(20)
    fix_core_properties(prs)
    prs.save(str(OUT / "unwrapped.pptx"))
    normalize_zip(OUT / "unwrapped.pptx")
    written(OUT / "unwrapped.pptx")


def symbol_bullets() -> None:
    """
    Bullets in the symbol fonts no phone has, set as PowerPoint and its designs set them: by
    the font's own code, or by that code in the private use area. The Circuit design's
    Wingdings 3 arrowhead (#48), PowerPoint's own arrow, one each from Wingdings, Wingdings 2,
    Symbol and Webdings, and LibreOffice's StarSymbol in the private use area, which it has no
    glyph for either. A second slide's body text takes a round Wingdings bullet from the
    master, as a design gives one. A third holds SmartArt, which is drawn from a part of its
    own, with a round Wingdings bullet and the Wingdings 3 arrowhead there.
    """
    from pptx import Presentation
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(6))
    box.name = "Bullets"
    frame = box.text_frame
    for n, (font, char, text) in enumerate((
        ("Wingdings 2", chr(0xF098), "A circle from Wingdings 2"),
        ("Wingdings 3", chr(0xF07D), "An arrowhead from Wingdings 3"),
        ("Wingdings", "\u00d8", "The arrow PowerPoint offers"),
        ("Wingdings", chr(0xF06C), "A round bullet from Wingdings"),
        ("Symbol", chr(0xF0B7), "A bullet from Symbol"),
        ("Webdings", "4", "A triangle from Webdings"),
        ("Wingdings 2", chr(0xF036), "A printer from Wingdings 2"),
        ("StarSymbol", chr(0xF06C), "A private use bullet from StarSymbol"),
    )):
        paragraph = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(24)
        props = paragraph._p.get_or_add_pPr()
        props.set("marL", "457200")
        props.set("indent", "-457200")
        props.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="{font}"/>'))
        props.append(parse_xml(f'<a:buChar {nsdecls("a")} char="{char}"/>'))

    level = prs.slide_master.element.find(qn("p:txStyles")).find(qn("p:bodyStyle")).find(qn("a:lvl1pPr"))
    level.find(qn("a:buFont")).set("typeface", "Wingdings")
    level.find(qn("a:buChar")).set("char", "l")
    designed = prs.slides.add_slide(prs.slide_layouts[1])
    designed.shapes.title.text = "From the design"
    body = designed.placeholders[1]
    body.name = "Designed"
    body.text = "A round bullet from the master"

    # SmartArt as PowerPoint saves it: a frame on the slide naming the diagram's data, layout,
    # style and colour parts, and the drawing PPTXjs draws it from, found by the slide's
    # relationships. Its shapes are named here, where PowerPoint leaves them unnamed, for the
    # tests to find them.
    dgm = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    smart = prs.slides.add_slide(prs.slide_layouts[6])
    ids = {}
    for kind, name, root, content in (
        ("diagramData", "data1", "dataModel", "diagramData"),
        ("diagramLayout", "layout1", "layoutDef", "diagramLayout"),
        ("diagramQuickStyle", "quickStyle1", "styleDef", "diagramStyle"),
        ("diagramColors", "colors1", "colorsDef", "diagramColors"),
    ):
        part = Part(PackURI(f"/ppt/diagrams/{name}.xml"),
                    f"application/vnd.openxmlformats-officedocument.drawingml.{content}+xml",
                    prs.part.package, f'<dgm:{root} xmlns:dgm="{dgm}"/>'.encode())
        ids[kind] = smart.part.relate_to(
            part, f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}")
    shapes = ""
    for n, (name, font, char, text) in enumerate((
        ("Round in SmartArt", "Wingdings", "l", "A round bullet in SmartArt"),
        ("Arrowhead in SmartArt", "Wingdings 3", chr(0xF07D), "An arrowhead in SmartArt"),
    )):
        shapes += (
            f'<dsp:sp modelId="{{00000000-0000-0000-0000-00000000000{n + 1}}}"><dsp:nvSpPr>'
            f'<dsp:cNvPr id="0" name="{name}"/><dsp:cNvSpPr/></dsp:nvSpPr><dsp:spPr><a:xfrm>'
            f'<a:off x="0" y="{n * 1371600}"/><a:ext cx="7315200" cy="1371600"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></dsp:spPr><dsp:txBody><a:bodyPr/>'
            f'<a:lstStyle/><a:p><a:pPr marL="457200" indent="-457200"><a:buFont typeface="{font}"/>'
            f'<a:buChar char="{char}"/></a:pPr><a:r><a:rPr lang="en-US" sz="2400"/><a:t>{text}</a:t>'
            '</a:r></a:p></dsp:txBody></dsp:sp>'
        )
    drawing = Part(PackURI("/ppt/diagrams/drawing1.xml"), "application/vnd.ms-office.drawingml.diagramDrawing+xml",
                   prs.part.package,
                   (f'<dsp:drawing xmlns:dsp="http://schemas.microsoft.com/office/drawing/2008/diagram" {nsdecls("a")}>'
                    '<dsp:spTree><dsp:nvGrpSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvGrpSpPr/></dsp:nvGrpSpPr>'
                    f'<dsp:grpSpPr/>{shapes}</dsp:spTree></dsp:drawing>').encode())
    smart.part.relate_to(drawing, "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing")
    smart.shapes._spTree.append(parse_xml(
        f'<p:graphicFrame {nsdecls("p", "a", "r")}><p:nvGraphicFramePr><p:cNvPr id="2" name="SmartArt"/>'
        '<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="914400" y="1371600"/><a:ext cx="7315200" cy="2743200"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{dgm}"><dgm:relIds xmlns:dgm="{dgm}" r:dm="{ids["diagramData"]}"'
        f' r:lo="{ids["diagramLayout"]}" r:qs="{ids["diagramQuickStyle"]}" r:cs="{ids["diagramColors"]}"/>'
        '</a:graphicData></a:graphic></p:graphicFrame>'))
    fix_core_properties(prs)
    prs.save(str(OUT / "symbol-bullets.pptx"))
    normalize_zip(OUT / "symbol-bullets.pptx")
    written(OUT / "symbol-bullets.pptx")


def placed_by_design() -> None:
    """
    Placeholders that give their shape an outline but leave out where it sits, as a deck
    converted from a PDF writes them (#48): a title that sits where the master's does, and a
    body where its layout's does. PPTXjs read that place from the shape alone, and threw.
    """
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]
    body = layout.placeholders.get(idx=1)
    body.left, body.top, body.width, body.height = Inches(5), Inches(2.5), Inches(4), Inches(3)
    slide = prs.slides.add_slide(layout)
    for shape, name, text in (
        (slide.shapes.title, "From the master", "Placed by the master"),
        (slide.placeholders[1], "From the layout", "Placed by the layout"),
    ):
        shape.name = name
        shape.text = text
        shape._element.spPr.append(parse_xml(f'<a:prstGeom {nsdecls("a")} prst="rect"><a:avLst/></a:prstGeom>'))
    fix_core_properties(prs)
    prs.save(str(OUT / "placed-by-design.pptx"))
    normalize_zip(OUT / "placed-by-design.pptx")
    written(OUT / "placed-by-design.pptx")


def unsized() -> None:
    """
    Text whose size nothing in the deck gives, not its run, its master's styles or the deck's
    defaults, beside a run that says 18 point (#48). PowerPoint draws both at 18 point, and
    PPTXjs drew the first at whatever size its paragraph had.
    """
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    prs = Presentation()
    for styles in (prs.slide_master.element.find(qn("p:txStyles")), prs.part._element.find(qn("p:defaultTextStyle"))):
        for props in styles.iter(qn("a:defRPr")):
            props.attrib.pop("sz", None)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for n, (name, text, size) in enumerate((("Unsized", "No size anywhere", None), ("Sized", "Eighteen point", Pt(18)))):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + 2 * n), Inches(8), Inches(1))
        box.name = name
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        if size:
            run.font.size = size
    fix_core_properties(prs)
    prs.save(str(OUT / "unsized.pptx"))
    normalize_zip(OUT / "unsized.pptx")
    written(OUT / "unsized.pptx")


def unreadable() -> None:
    """
    Things PPTXjs throws on, which stopped every slide (#48): a chart whose part is missing
    from the file, which gave the reporter's "reading 'c:chartSpace'", and a picture with no
    image on a layout, which PPTXjs draws behind each slide that uses it. The layout's text
    must still be drawn, and a slide with neither sits between them. The layout also has an
    empty picture placeholder that PPTXjs throws on, which PowerPoint shows on no slide, and
    the slide between has two shaded boxes PPTXjs throws on (a gradient with no colours,
    as in POI's 63200.pptx), one with text and one with none.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for xml in (
        '<p:pic {ns}><p:nvPicPr><p:cNvPr id="50" name="No image"/><p:cNvPicPr/><p:nvPr userDrawn="1"/></p:nvPicPr>'
        '<p:blipFill/><p:spPr><a:xfrm><a:off x="457200" y="457200"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>',
        '<p:sp {ns}><p:nvSpPr><p:cNvPr id="51" name="Layout text"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/>'
        '</p:nvSpPr><p:spPr><a:xfrm><a:off x="457200" y="5943600"/><a:ext cx="4572000" cy="457200"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>'
        '<a:p><a:r><a:rPr lang="en-US" sz="1800"/><a:t>Drawn by the layout</a:t></a:r></a:p></p:txBody></p:sp>',
    ):
        blank.shapes._spTree.append(parse_xml(xml.format(ns=nsdecls("p", "a"))))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A chart Gander cannot read"
    data = CategoryChartData()
    data.categories = ["North", "South"]
    data.add_series("Visits", (3, 5))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), data)
    slide.shapes.add_textbox(Inches(5.5), Inches(2), Inches(3), Inches(1)).text_frame.text = "Beside the chart"
    after = prs.slides.add_slide(prs.slide_layouts[5])
    after.shapes.title.text = "The slide after it"
    for xml in (
        '<p:sp {ns}><p:nvSpPr><p:cNvPr id="60" name="Shaded, no text"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="914400" y="1828800"/><a:ext cx="1828800" cy="914400"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:gradFill/></p:spPr></p:sp>',
        '<p:sp {ns}><p:nvSpPr><p:cNvPr id="61" name="Shaded, with text"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="4572000" y="1828800"/><a:ext cx="2743200" cy="914400"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:gradFill/></p:spPr><p:txBody><a:bodyPr/>'
        '<a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1800"/><a:t>Words in a shaded box</a:t></a:r></a:p>'
        '</p:txBody></p:sp>',
    ):
        after.shapes._spTree.append(parse_xml(xml.format(ns=nsdecls("p", "a"))))
    box = prs.slides.add_slide(blank).shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    box.text_frame.text = "On a layout with a broken picture"
    # Added after the slide, which would otherwise be given a copy of it
    blank.shapes._spTree.append(parse_xml(
        '<p:pic {ns}><p:nvPicPr><p:cNvPr id="52" name="Picture placeholder"/><p:cNvPicPr/>'
        '<p:nvPr><p:ph type="pic" idx="13"/></p:nvPr></p:nvPicPr><p:blipFill/><p:spPr><a:xfrm>'
        '<a:off x="5486400" y="457200"/><a:ext cx="914400" cy="914400"/></a:xfrm></p:spPr></p:pic>'
        .format(ns=nsdecls("p", "a"))))
    fix_core_properties(prs)
    path = OUT / "unreadable.pptx"
    prs.save(str(path))

    # The chart's part goes, with its workbook, and the slide's relationship to it stays
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()
                 if not i.filename.startswith(("ppt/charts/", "ppt/embeddings/"))}
    items["[Content_Types].xml"] = re.sub(rb'<Override PartName="/ppt/charts/[^"]*"[^>]*/>', b"", items["[Content_Types].xml"])
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)
    normalize_zip(path)
    written(path)


def charted() -> None:
    """
    A chart PPTXjs reads and hands to nv.d3, which draws it once every slide is built, with
    text beside it and a slide after it. test_pptx makes the drawing throw, as an odd chart
    could inside nv.d3, which put the error card up in place of every slide (#48).
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A chart nv.d3 draws"
    data = CategoryChartData()
    data.categories = ["North", "South"]
    data.add_series("Visits", (3, 5))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), data)
    slide.shapes.add_textbox(Inches(5.5), Inches(2), Inches(3), Inches(1)).text_frame.text = "Beside the chart"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "The slide after it"
    fix_core_properties(prs)
    path = OUT / "charted.pptx"
    prs.save(str(path))
    normalize_zip(path)
    written(path)


def undrawn() -> None:
    """
    Pictures and a chart that PPTXjs puts on the slide and that never draw, though nothing
    throws: a Windows metafile, which no browser draws, beside a PNG that does and a PNG cut
    short, and a doughnut chart, a type PPTXjs does not draw. pptx.js marks each where it is
    (#48).
    """
    from PIL import Image
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    # A placeable metafile of one rectangle 3 by 2 inches, as a drawing pasted from another
    # program is kept: its header, with a checksum of the words before it, then the records
    inch, right, bottom = 1440, 4320, 2880
    head = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, right, bottom, inch, 0)
    checksum = 0
    for (word,) in struct.iter_unpack("<H", head):
        checksum ^= word
    records = (struct.pack("<IHhh", 5, 0x020B, 0, 0) + struct.pack("<IHhh", 5, 0x020C, bottom, right)
               + struct.pack("<IHhhhh", 7, 0x041B, bottom, right, 0, 0) + struct.pack("<IH", 3, 0))
    metafile = (head + struct.pack("<H", checksum)
                + struct.pack("<HHHIHIH", 1, 9, 0x0300, (18 + len(records)) // 2, 0, 7, 0) + records)
    png, damaged = io.BytesIO(), io.BytesIO()
    Image.new("RGB", (60, 40), (70, 130, 180)).save(png, "PNG")
    Image.new("RGB", (60, 40), (180, 70, 70)).save(damaged, "PNG")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A drawing in Windows' own format"
    slide.shapes.add_picture(io.BytesIO(metafile), Inches(1), Inches(2), Inches(3), Inches(2))
    slide.shapes.add_picture(png, Inches(5.5), Inches(2), Inches(3), Inches(2))
    slide.shapes.add_picture(io.BytesIO(damaged.getvalue()), Inches(5.5), Inches(4.5), Inches(3), Inches(2))
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A doughnut chart"
    data = CategoryChartData()
    data.categories = ["North", "South"]
    data.add_series("Visits", (3, 5))
    slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(1), Inches(2), Inches(4), Inches(3), data)
    slide.shapes.add_textbox(Inches(5.5), Inches(2), Inches(3), Inches(1)).text_frame.text = "Beside the chart"
    fix_core_properties(prs)
    path = OUT / "undrawn.pptx"
    prs.save(str(path))

    # The second PNG ends after its header, as a file cut off in a copy would
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}
    for name, data in items.items():
        if data == damaged.getvalue():
            items[name] = data[:33]
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)
    normalize_zip(path)
    written(path)


def grounds() -> None:
    """
    Issue #53: a master whose background is a pattern, as the reported deck's is, which
    PPTXjs cannot draw. The first slide takes it, the second says it has no background,
    and the third has its own dark blue. PPTXjs drew the first two see-through.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_PATTERN, MSO_THEME_COLOR
    from pptx.util import Inches, Pt

    prs = Presentation()
    pattern = prs.slide_master.background.fill
    pattern.patterned()
    pattern.pattern = MSO_PATTERN.LIGHT_HORIZONTAL
    pattern.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_2
    pattern.back_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
    for text, ink in (("On the master's pattern", "1F3864"), ("No background", "1F3864"),
                      ("Its own dark blue", "FFFFFF")):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor.from_string(ink)
    prs.slides[1].background.fill.background()
    own = prs.slides[2].background.fill
    own.solid()
    own.fore_color.rgb = RGBColor.from_string("1F3864")
    fix_core_properties(prs)
    target = OUT / "grounds.pptx"
    prs.save(str(target))
    normalize_zip(target)
    written(target)


# What [Content_Types].xml declares each format's main part to be. The rest of a
# package is the same across a family, so this one line is all that tells a
# template, a slide show or a macro-enabled file from its format, and a reader
# that looks its main part up by type is the one that would refuse it.
MAIN_PARTS = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "docm": "application/vnd.ms-word.document.macroEnabled.main+xml",
    "dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "xltx": "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    "xlsm": "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
    "ppsx": "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
    "pptm": "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    "potx": "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
}


def retyped(source: str, ext: str) -> None:
    """Writes the fixture [source] again as .[ext], its main part declared as that format's.

    Nothing else changes. A .docm or a .pptm with macros in it carries them as a
    part of their own as well, which no viewer here reads, so none is invented.
    """
    stem, was = source.rsplit(".", 1)
    old, new = MAIN_PARTS[was].encode(), MAIN_PARTS[ext].encode()
    target = OUT / f"{stem}.{ext}"
    with zipfile.ZipFile(OUT / source) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(target, "w") as z:
        for name, data in items:
            if name == "[Content_Types].xml":
                assert data.count(old) == 1, f"{source} does not declare its main part once"
                data = data.replace(old, new)
            z.writestr(name, data)
    normalize_zip(target)
    written(target)


def relatives() -> None:
    """One of each Office format's relatives, made from the fixtures above."""
    for ext in ("docm", "dotx"):
        retyped("report.docx", ext)
    for ext in ("xltx", "xlsm"):
        retyped("budget.xlsx", ext)
    for ext in ("ppsx", "pptm", "potx"):
        retyped("deck.pptx", ext)


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------

def texts() -> None:
    md = """# Willowmere site notes

A heading, a [link](https://example.invalid/notes), and a list:

- first
- second

<script>window.__xss = 1;</script>

<img src="x" onerror="window.__xss = 2;">

Text after the injected markup, so the sanitiser can be seen to have kept it.
"""
    (OUT / "notes.md").write_text(md, encoding="utf-8")
    written(OUT / "notes.md")

    plain = (
        "Plain text, opened by the text viewer.\n"
        "A second line so the newline handling is visible.\n"
        "An accented character: café.\n"
    )
    (OUT / "plain.txt").write_text(plain, encoding="utf-8")
    written(OUT / "plain.txt")

    # Issue #32. A .json file as they nearly always arrive: no spacing at all,
    # the whole document on one line. The id is 2**53 + 1, the smallest whole
    # number a double cannot hold, so a viewer that reformatted by parsing and
    # restringifying would show 9007199254740992 here and be caught. The text
    # carries an accent and an emoji, which is a surrogate pair, because both
    # sit inside a string the formatter has to copy through untouched.
    snapshot = (
        '{"id":9007199254740993,"app":"com.example.reader",'
        '"screen":{"width":1080,"height":2376},"landscape":false,"tags":[],'
        '"nodes":[{"id":0,"parent":-1,"name":"android.widget.FrameLayout",'
        '"text":null,"visible":true},{"id":1,"parent":0,'
        '"name":"android.widget.TextView","text":"café \U0001f600",'
        '"visible":true}],"notes":{}}'
    )
    (OUT / "snapshot.json").write_text(snapshot, encoding="utf-8")
    written(OUT / "snapshot.json")

    # Byte order marks. app.js sniffs these three bytes and picks the decoder;
    # the decoder strips the mark itself, so neither file should show one.
    marked = "Byte order marked text, decoded by the mark alone.\n"
    (OUT / "utf16le.txt").write_bytes(b"\xff\xfe" + marked.encode("utf-16-le"))
    written(OUT / "utf16le.txt")
    (OUT / "utf16be.txt").write_bytes(b"\xfe\xff" + marked.encode("utf-16-be"))
    written(OUT / "utf16be.txt")

    # An extension nothing claims, over bytes that read as text, so the app opens it as text.
    (OUT / "unknown.xyz").write_bytes(bytes(range(32, 127)) * 4 + b"\n")
    written(OUT / "unknown.xyz")

    # The same over bytes that do not, a zero among them, so the app keeps its card.
    (OUT / "opaque.xyz").write_bytes(bytes(range(256)) * 2)
    written(OUT / "opaque.xyz")

    # Text under other names, which the text viewer shows as it is: subtitles
    # as SubRip and as WebVTT, a playlist of tracks that are not here, as one
    # always arrives, and the notes that come with a download.
    captions = [
        ("00:00:01", "000", "00:00:03", "500", "The ferry leaves Willowmere at nine."),
        ("00:00:04", "000", "00:00:06", "250", "Bring the survey maps, and a coat."),
    ]
    srt = "\n".join(
        f"{n}\n{a},{ams} --> {b},{bms}\n{line}\n"
        for n, (a, ams, b, bms, line) in enumerate(captions, start=1)
    )
    (OUT / "captions.srt").write_text(srt, encoding="utf-8")
    written(OUT / "captions.srt")
    vtt = "WEBVTT\n\n" + "\n".join(
        f"{a}.{ams} --> {b}.{bms}\n{line}\n" for a, ams, b, bms, line in captions
    )
    (OUT / "captions.vtt").write_text(vtt, encoding="utf-8")
    written(OUT / "captions.vtt")

    playlist = (
        "#EXTM3U\n"
        "#EXTINF:187,Willowmere Field Recordings - Rain on the Survey Hut\n"
        "Music/Willowmere/01 Rain on the Survey Hut.mp3\n"
        "#EXTINF:204,Willowmere Field Recordings - The Drainage Ditch at Dusk\n"
        "Music/Willowmere/02 The Drainage Ditch at Dusk.mp3\n"
    )
    (OUT / "playlist.m3u").write_text(playlist, encoding="utf-8")
    written(OUT / "playlist.m3u")

    nfo = (
        "  WILLOWMERE FIELD RECORDINGS\n"
        "  ---------------------------\n"
        "\n"
        "  Recorded ....: January 2026\n"
        "  Tracks ......: 2\n"
        "  Format ......: MP3, 192 kbps\n"
        "\n"
        "  Two recordings made up for a test, so that the text\n"
        "  viewer has something to show. None of it is real.\n"
    )
    (OUT / "release.nfo").write_text(nfo, encoding="utf-8")
    written(OUT / "release.nfo")

    verilog = (
        "// Willowmere pump controller, made up for a test: the pump runs\n"
        "// while the tank reads low.\n"
        "module pump(input wire clk, input wire tank_low, output reg running);\n"
        "  always @(posedge clk)\n"
        "    running <= tank_low;\n"
        "endmodule\n"
    )
    (OUT / "pump.v").write_text(verilog, encoding="utf-8")
    written(OUT / "pump.v")



# ---------------------------------------------------------------------------
# Images pdf.js decodes in WebAssembly
# ---------------------------------------------------------------------------

# Issue #24. Since pdf.js 4 three image encodings are decoded by wasm modules
# the worker fetches on demand rather than by JavaScript in the bundle, and the
# only way to say where those modules are is the wasmUrl option. Gander shipped
# neither the binaries nor the option until 1.17, so every image in these three
# encodings was dropped: the worker warns to a console nobody reads, returns
# nothing, and the page renders with a hole where the picture goes. No error, no
# placeholder, nothing that looks like a failure rather than a layout.
#
# One module, jbig2.wasm, serves both JBIG2 and CCITT fax. CCITT is the ordinary
# compression for a scanned black and white page and by far the commoner of the
# two, and it had been broken the whole time without anyone reporting it, which
# is the argument for testing all three rather than the one that was reported.
#
# Written by hand rather than through ReportLab because the point of each file
# is its /Filter, and a library that re-encodes to something it prefers would
# quietly test nothing.


def _raw_image_pdf(path: Path, w: int, h: int, entries: bytes, data: bytes) -> None:
    """A one-page PDF whose only content is a single image XObject, undecoded."""
    content = b"q 480 0 0 200 20 30 cm /Im0 Do Q"
    objs = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 520 260] "
            b"/Resources << /XObject << /Im0 5 0 R >> >> /Contents 4 0 R >>"),
        4: b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        5: (b"<< /Type /XObject /Subtype /Image /Width %d /Height %d "
            b"%s /Length %d >>\nstream\n" % (w, h, entries, len(data))
            + data + b"\nendstream"),
    }
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for n in sorted(objs):
        offsets[n] = len(out)
        out += b"%d 0 obj\n" % n + objs[n] + b"\nendobj\n"
    start = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for n in sorted(objs):
        out += b"%010d 00000 n \n" % offsets[n]
    out += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objs) + 1, start))
    path.write_bytes(bytes(out))


def _bitonal(w: int, h: int):
    """Black on white, with a circle and two text runs, so a partial decode shows."""
    from PIL import Image, ImageDraw
    im = Image.new("1", (w, h), 1)
    d = ImageDraw.Draw(im)
    d.rectangle([10, 10, w - 11, h - 11], outline=0, width=3)
    d.text((40, 60), "SCANNED PAGE", fill=0)
    d.text((40, 90), "this line proves the decoder ran", fill=0)
    d.ellipse([w - 150, 40, w - 50, 140], outline=0, width=4)
    return im


def _group4(im) -> bytes:
    """T.6 data, taken out of a Group 4 TIFF because Pillow will write one."""
    from PIL import Image
    buf = io.BytesIO()
    im.save(buf, format="TIFF", compression="group4")
    buf.seek(0)
    tiff = Image.open(buf)
    raw = buf.getvalue()
    return b"".join(raw[o:o + c] for o, c in
                    zip(tiff.tag_v2[273], tiff.tag_v2[279]))


def wasm_decoded_images() -> None:
    from PIL import Image, ImageDraw

    W, H = 480, 200

    # jpx.pdf: JPEG 2000, which is what issue #24's file was made of. A
    # photograph rather than a diagram, because JPEG 2000 is a photographic
    # codec and a flat drawing would compress to something unrepresentative.
    photo = Image.new("RGB", (W, H))
    px = photo.load()
    for y in range(H):
        for x in range(W):
            px[x, y] = ((x * 255) // W, (y * 255) // H, ((x + y) * 127) // (W + H))
    d = ImageDraw.Draw(photo)
    d.ellipse([W - 150, 40, W - 50, 140], fill=(250, 250, 40))
    d.text((40, 90), "JPEG 2000", fill=(255, 255, 255))
    buf = io.BytesIO()
    photo.save(buf, format="JPEG2000", irreversible=True, quality_layers=[40])
    # No /ColorSpace: for JPXDecode the codestream carries it, and naming one
    # here would let a reader that ignored the image still look correct.
    _raw_image_pdf(OUT / "jpx.pdf", W, H, b"/Filter /JPXDecode", buf.getvalue())
    written(OUT / "jpx.pdf")

    bitonal = _bitonal(W, H)
    mmr = _group4(bitonal)

    # jbig2.pdf: an embedded JBIG2 stream, which is a bare segment sequence with
    # no file header. A generic region with MMR=1 carries plain T.6 data, so the
    # Group 4 bytes above can be reused and the fixture needs no JBIG2 encoder
    # on the machine generating it.
    page_info = struct.pack(">IIII", W, H, 0, 0) + bytes([0x01]) + struct.pack(">H", 0)
    region = (struct.pack(">IIII", W, H, 0, 0)   # region position and size
              + bytes([0x00])                    # combine into the page by OR
              + bytes([0x01])                    # MMR = 1, so no arithmetic coder
              + mmr)

    def segment(number: int, kind: int, data: bytes) -> bytes:
        return (struct.pack(">I", number)
                + bytes([kind & 0x3F])   # one-byte page association
                + bytes([0x00])          # refers to no other segment
                + bytes([0x01])          # page 1
                + struct.pack(">I", len(data))
                + data)

    jb2 = segment(0, 48, page_info) + segment(1, 39, region)
    _raw_image_pdf(OUT / "jbig2.pdf", W, H,
                   b"/ColorSpace /DeviceGray /BitsPerComponent 1 /Filter /JBIG2Decode",
                   jb2)
    written(OUT / "jbig2.pdf")

    # ccitt.pdf: the same bitmap as plain Group 4, which is what a scanner or a
    # fax produces and what jbig2.wasm turns out to decode as well.
    _raw_image_pdf(OUT / "ccitt.pdf", W, H,
                   b"/ColorSpace /DeviceGray /BitsPerComponent 1 "
                   b"/Filter /CCITTFaxDecode "
                   b"/DecodeParms << /K -1 /Columns %d /Rows %d /BlackIs1 true >>"
                   % (W, H),
                   mmr)
    written(OUT / "ccitt.pdf")


# ---------------------------------------------------------------------------
# Images and audio
# ---------------------------------------------------------------------------

# The six EXIF orientations Thumbs.exifRotation maps to a rotation, plus the
# two flips it folds into 90 and 270.
EXIF_ORIENTATIONS = {
    1: 0,    # normal
    3: 180,  # rotate 180
    6: 90,   # rotate 90
    8: 270,  # rotate 270
    5: 90,   # transpose
    7: 270,  # transverse
}


def images() -> None:
    from PIL import Image, ImageDraw

    def asymmetric(w=120, h=80):
        """A frame with one filled corner, so a rotation is visible."""
        img = Image.new("RGB", (w, h), (245, 245, 245))
        d = ImageDraw.Draw(img)
        d.rectangle([0, 0, w - 1, h - 1], outline=(30, 30, 30), width=2)
        d.rectangle([4, 4, 34, 24], fill=(178, 45, 24))
        return img

    for orientation in sorted(EXIF_ORIENTATIONS):
        path = OUT / f"exif-{orientation}.jpg"
        exif = Image.Exif()
        exif[0x0112] = orientation
        asymmetric().save(path, "JPEG", quality=88, exif=exif)
        written(path)

    asymmetric(64, 64).save(OUT / "tiny.png", "PNG", optimize=True)
    written(OUT / "tiny.png")

    frames = []
    for shift in range(4):
        img = Image.new("P", (48, 48), 0)
        d = ImageDraw.Draw(img)
        d.rectangle([shift * 8, 8, shift * 8 + 16, 24], fill=1)
        img.putpalette([245, 245, 245, 178, 45, 24] + [0] * 762)
        frames.append(img)
    frames[0].save(
        OUT / "anim.gif", save_all=True, append_images=frames[1:],
        duration=120, loop=0,
    )
    written(OUT / "anim.gif")

    # One picture that moves, as an animated WebP and as an animated PNG, a whole colour to a
    # frame so a test can see it change, and a WebP that keeps still (#49)
    moving = [Image.new("RGB", (48, 48), c) for c in ((178, 45, 24), (40, 140, 70), (40, 70, 190))]
    for path in (OUT / "moving.webp", OUT / "moving.png"):
        moving[0].save(path, save_all=True, append_images=moving[1:], duration=300, loop=0, lossless=True)
        written(path)
    asymmetric(64, 64).save(OUT / "still.webp", "WEBP", lossless=True)
    written(OUT / "still.webp")

    (OUT / "icon.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80" '
        'viewBox="0 0 120 80" role="img" aria-label="A red square in a frame">\n'
        '  <rect width="120" height="80" fill="#f5f5f5" stroke="#1e1e1e" '
        'stroke-width="2"/>\n'
        '  <rect x="8" y="8" width="30" height="20" fill="#b22d18"/>\n'
        "</svg>\n",
        encoding="utf-8",
    )
    written(OUT / "icon.svg")


def audio() -> None:
    rate, seconds, freq = 8000, 1, 440.0
    frames = bytearray()
    for i in range(rate * seconds):
        # A quiet sine, so a device test that actually plays it is bearable
        frames += struct.pack("<h", int(6000 * sin(2 * pi * freq * i / rate)))
    with wave.open(str(OUT / "tone.wav"), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(bytes(frames))
    written(OUT / "tone.wav")


def photo_formats() -> None:
    """One photo as WebP, BMP and HEIC, the photo formats the tiling view leaves to Android.

    160 by 96, a size no other picture here has, so a device test that reads the decoded
    size back knows this file was read. The WebP is lossy, as a camera or a website saves
    one. HEIC needs pillow-heif, since Pillow writes no HEVC.
    """
    from PIL import Image, ImageDraw
    from pillow_heif import from_pillow

    img = Image.new("RGB", (160, 96), (245, 245, 245))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 159, 95], outline=(30, 30, 30), width=2)
    d.rectangle([6, 6, 54, 36], fill=(178, 45, 24))
    img.save(OUT / "photo.webp", "WEBP", quality=80, method=6)
    written(OUT / "photo.webp")
    img.save(OUT / "photo.bmp", "BMP")
    written(OUT / "photo.bmp")
    from_pillow(img).save(OUT / "photo.heic", quality=80)
    written(OUT / "photo.heic")


# ---------------------------------------------------------------------------
# Zips, issue #30
# ---------------------------------------------------------------------------

# Written by hand, like the wasm PDFs above and for the same reason: the point of
# most of these is a byte Python's zipfile will not write. It sets the UTF-8 flag
# on every name that is not ASCII, and the names Gander has to decode are exactly
# the ones a Windows machine writes in its own code page with that flag clear.

ZIP_DOS_DATE = ((ZIP_DATE[0] - 1980) << 9) | (ZIP_DATE[1] << 5) | ZIP_DATE[2]
HOST_DOS, HOST_UNIX = 0, 3
STORED, DEFLATED, DEFLATE64 = 0, 8, 9
AES = 99


class Member:
    """One entry. A method other than stored, deflated or Deflate64 writes data as given.

    zipcrypto or aes, a password, encrypts it: aes as (strength, version), strength 1 to 3
    for 128 to 256 bit keys and version 1 or 2 for AE-1 or AE-2.
    """

    def __init__(self, name: bytes, data: bytes = b"", *, method=DEFLATED, flags=0,
                 host=HOST_UNIX, extra=b"", descriptor=False, directory=False, time=0,
                 zipcrypto: bytes = None, aes=None, password: bytes = None):
        self.name, self.data, self.method = name, data, method
        self.flags = flags | (0x08 if descriptor else 0) | (0x01 if zipcrypto or aes else 0)
        self.host, self.extra = host, extra
        self.descriptor, self.directory, self.time = descriptor, directory, time
        self.zipcrypto, self.aes, self.password = zipcrypto, aes, password


def _extra(field_id: int, data: bytes) -> bytes:
    return struct.pack("<HH", field_id, len(data)) + data


# PKWARE's original encryption, APPNOTE 6.1. Twelve bytes in front, the last of them the
# top byte of the CRC, or of the time when the CRC trails the data; the eleven before it
# would be random, and are a hash of the name here so the output never changes.

def _crc_byte(crc: int, b: int) -> int:
    return zlib.crc32(bytes([b]), crc ^ 0xFFFFFFFF) ^ 0xFFFFFFFF


class _ZipCrypto:
    def __init__(self, password: bytes):
        self.k = [0x12345678, 0x23456789, 0x34567890]
        for b in password:
            self._update(b)

    def _update(self, b: int) -> None:
        k0 = _crc_byte(self.k[0], b)
        k1 = ((self.k[1] + (k0 & 0xFF)) * 134775813 + 1) & 0xFFFFFFFF
        self.k = [k0, k1, _crc_byte(self.k[2], k1 >> 24)]

    def encrypt(self, data: bytes) -> bytes:
        out = bytearray()
        for b in data:
            t = (self.k[2] | 2) & 0xFFFF
            out.append(b ^ (((t * (t ^ 1)) >> 8) & 0xFF))
            self._update(b)
        return bytes(out)


def zipcrypto(body: bytes, password: bytes, check: int, name: bytes) -> bytes:
    keys = _ZipCrypto(password)
    header = hashlib.sha256(b"header" + name).digest()[:11] + bytes([check])
    return keys.encrypt(header) + keys.encrypt(body)


# WinZip's AES, from its published specification: PBKDF2-HMAC-SHA1 over the password and a
# salt, a thousand rounds, into the AES key, an HMAC key and a two byte verifier; AES in
# counter mode counting up from one little-endian; the first ten bytes of HMAC-SHA1 over what
# was encrypted, after it. The salt is a hash of the name, again so nothing changes.

def winzip_aes(body: bytes, password: bytes, strength: int, name: bytes) -> bytes:
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    size = {1: 16, 2: 24, 3: 32}[strength]
    salt = hashlib.sha256(b"salt" + name).digest()[:size // 2]
    keys = hashlib.pbkdf2_hmac("sha1", password, salt, 1000, 2 * size + 2)
    aes = Cipher(algorithms.AES(keys[:size]), modes.ECB()).encryptor()
    out = bytearray()
    for n, at in enumerate(range(0, len(body), 16), start=1):
        stream = aes.update(n.to_bytes(16, "little"))
        out += bytes(a ^ b for a, b in zip(body[at:at + 16], stream))
    code = hmac.new(keys[size:2 * size], bytes(out), hashlib.sha1).digest()[:10]
    return salt + keys[2 * size:] + bytes(out) + code


# Deflate64, which zlib cannot write. Enough of an encoder to put every part of the format
# in one stream: a stored block, a fixed one and dynamic ones, matches reaching into the
# second 32 KB of the window, and lengths past 258, which only Deflate64's last length code
# can carry. Greedy matching over the whole 64 KB window, nothing cleverer.

_LEN_BASE = [3, 4, 5, 6, 7, 8, 9, 10, 11, 13, 15, 17, 19, 23, 27, 31, 35, 43, 51, 59, 67, 83,
             99, 115, 131, 163, 195, 227, 3]
_LEN_EXTRA = [0, 0, 0, 0, 0, 0, 0, 0, 1, 1, 1, 1, 2, 2, 2, 2, 3, 3, 3, 3, 4, 4, 4, 4, 5, 5, 5,
              5, 16]
_DIST_BASE = [1, 2, 3, 4, 5, 7, 9, 13, 17, 25, 33, 49, 65, 97, 129, 193, 257, 385, 513, 769,
              1025, 1537, 2049, 3073, 4097, 6145, 8193, 12289, 16385, 24577, 32769, 49153]
_DIST_EXTRA = [0, 0, 0, 0, 1, 1, 2, 2, 3, 3, 4, 4, 5, 5, 6, 6, 7, 7, 8, 8, 9, 9, 10, 10, 11,
               11, 12, 12, 13, 13, 14, 14]
_ORDER = [16, 17, 18, 0, 8, 7, 9, 6, 10, 5, 11, 4, 12, 3, 13, 2, 14, 1, 15]


class _Bits:
    def __init__(self):
        self.out, self.acc, self.n = bytearray(), 0, 0

    def put(self, value: int, count: int) -> None:
        self.acc |= value << self.n
        self.n += count
        while self.n >= 8:
            self.out.append(self.acc & 0xFF)
            self.acc >>= 8
            self.n -= 8

    def code(self, code: int, length: int) -> None:
        self.put(int(format(code, f"0{length}b")[::-1], 2), length)

    def align(self) -> None:
        if self.n:
            self.out.append(self.acc & 0xFF)
        self.acc = self.n = 0


def _lengths(freqs, limit):
    used = [(f, s) for s, f in enumerate(freqs) if f]
    lengths = [0] * len(freqs)
    if not used:
        # A block of nothing but literals has no distance code at all, which is allowed
        return lengths
    if len(used) == 1:
        # One code alone is incomplete; a second, never used, completes it
        lengths[used[0][1]] = 1
        lengths[0 if used[0][1] else 1] = 1
        return lengths
    while True:
        heap = [(f, i, (s,)) for i, (f, s) in enumerate(used)]
        heapq.heapify(heap)
        depth, tie = {s: 0 for _, s in used}, len(heap)
        while len(heap) > 1:
            f1, _, a = heapq.heappop(heap)
            f2, _, b = heapq.heappop(heap)
            for s in a + b:
                depth[s] += 1
            heapq.heappush(heap, (f1 + f2, tie, a + b))
            tie += 1
        if max(depth.values()) <= limit:
            for s, d in depth.items():
                lengths[s] = d
            return lengths
        used = [((f + 1) // 2, s) for f, s in used]


def _canonical(lengths):
    count = [0] * 16
    for n in lengths:
        if n:
            count[n] += 1
    code, first = 0, [0] * 16
    for n in range(1, 16):
        code = (code + count[n - 1]) << 1
        first[n] = code
    codes = []
    for n in lengths:
        codes.append(first[n])
        if n:
            first[n] += 1
    return codes


def _length_code(n):
    if n > 258:
        return 285, n - 3, 16
    i = max(i for i in range(28) if _LEN_BASE[i] <= n)
    return 257 + i, n - _LEN_BASE[i], _LEN_EXTRA[i]


def _distance_code(d):
    i = max(i for i in range(32) if _DIST_BASE[i] <= d)
    return i, d - _DIST_BASE[i], _DIST_EXTRA[i]


def _matches(data: bytes, start: int):
    """Greedy LZ77 over a 64 KB window, from start, reaching back before it."""
    table, tokens, i = {}, [], 0
    for j in range(max(0, start - 65536), start):
        table.setdefault(data[j:j + 3], []).append(j)
    i = start
    while i < len(data):
        best, where = 0, 0
        for j in reversed(table.get(data[i:i + 3], [])[-32:]):
            if i - j > 65536:
                break
            n = 0
            while i + n < len(data) and n < 65538 and data[j + n] == data[i + n]:
                n += 1
            if n > best:
                best, where = n, i - j
        step = best if best >= 3 else 1
        for j in range(i, min(i + step, len(data) - 2)):
            table.setdefault(data[j:j + 3], []).append(j)
        tokens.append((best, where) if best >= 3 else data[i])
        i += step
    return tokens


def _block(bits: _Bits, tokens, last: bool, dynamic: bool) -> None:
    if dynamic:
        lit, dist = [0] * 286, [0] * 32
        lit[256] = 1
        for t in tokens:
            if isinstance(t, int):
                lit[t] += 1
            else:
                lit[_length_code(t[0])[0]] += 1
                dist[_distance_code(t[1])[0]] += 1
        lit_lengths, dist_lengths = _lengths(lit, 15), _lengths(dist, 15)
        hlit = max(257, max(s for s, n in enumerate(lit_lengths) if n) + 1)
        hdist = max(1, max((s for s, n in enumerate(dist_lengths) if n), default=0) + 1)
        seq, all_lengths, i = [], lit_lengths[:hlit] + dist_lengths[:hdist], 0
        while i < len(all_lengths):
            n = all_lengths[i]
            run = 1
            while i + run < len(all_lengths) and all_lengths[i + run] == n:
                run += 1
            if n == 0 and run >= 3:
                r = min(run, 138)
                seq.append((18, r - 11, 7) if r >= 11 else (17, r - 3, 3))
                i += r
            elif n and run >= 4:
                r = min(run - 1, 6)
                seq += [(n, 0, 0), (16, r - 3, 2)]
                i += 1 + r
            else:
                seq.append((n, 0, 0))
                i += 1
        cl = [0] * 19
        for sym, _, _ in seq:
            cl[sym] += 1
        cl_lengths = _lengths(cl, 7)
        hclen = 19
        while hclen > 4 and cl_lengths[_ORDER[hclen - 1]] == 0:
            hclen -= 1
        bits.put(1 if last else 0, 1)
        bits.put(2, 2)
        bits.put(hlit - 257, 5)
        bits.put(hdist - 1, 5)
        bits.put(hclen - 4, 4)
        for i in range(hclen):
            bits.put(cl_lengths[_ORDER[i]], 3)
        cl_codes = _canonical(cl_lengths)
        for sym, value, extra in seq:
            bits.code(cl_codes[sym], cl_lengths[sym])
            if extra:
                bits.put(value, extra)
    else:
        lit_lengths = [8] * 144 + [9] * 112 + [7] * 24 + [8] * 8
        dist_lengths = [5] * 32
        bits.put(1 if last else 0, 1)
        bits.put(1, 2)
    lit_codes, dist_codes = _canonical(lit_lengths), _canonical(dist_lengths)
    for t in tokens:
        if isinstance(t, int):
            bits.code(lit_codes[t], lit_lengths[t])
        else:
            sym, value, extra = _length_code(t[0])
            bits.code(lit_codes[sym], lit_lengths[sym])
            bits.put(value, extra)
            sym, value, extra = _distance_code(t[1])
            bits.code(dist_codes[sym], dist_lengths[sym])
            bits.put(value, extra)
    bits.code(lit_codes[256], lit_lengths[256])


def deflate64(data: bytes) -> bytes:
    """A stored block, then a fixed one, then two dynamic ones."""
    bits = _Bits()
    stored = data[:min(len(data), 16 * 1024)]
    bits.put(0, 1)
    bits.put(0, 2)
    bits.align()
    bits.out += struct.pack("<HH", len(stored), len(stored) ^ 0xFFFF) + stored
    tokens = _matches(data, len(stored))
    fixed, rest = tokens[:500], tokens[500:]
    half = len(rest) // 2
    _block(bits, fixed, last=False, dynamic=False)
    _block(bits, rest[:half], last=False, dynamic=True)
    _block(bits, rest[half:], last=True, dynamic=True)
    bits.align()
    return bytes(bits.out)


def zip_bytes(members, *, zip64=False, comment=b"") -> bytes:
    out, central = bytearray(), bytearray()
    for m in members:
        crc = zlib.crc32(m.data) & 0xFFFFFFFF
        if m.method == DEFLATED:
            packer = zlib.compressobj(9, zlib.DEFLATED, -15)
            body = packer.compress(m.data) + packer.flush()
        elif m.method == DEFLATE64:
            body = deflate64(m.data)
        else:
            body = m.data
        method, aes_extra = m.method, b""
        if m.zipcrypto:
            check = (m.time >> 8) if m.descriptor else (crc >> 24)
            body = zipcrypto(body, m.zipcrypto, check, m.name)
        elif m.aes:
            strength, version = m.aes
            body = winzip_aes(body, m.password, strength, m.name)
            aes_extra = _extra(0x9901, struct.pack("<H2sBH", version, b"AE", strength, m.method))
            method = AES
            if version == 2:
                crc = 0
        if method == AES:
            needed = 51
        elif method == DEFLATE64:
            needed = 21
        else:
            needed = 63 if method not in (STORED, DEFLATED) else (45 if zip64 else 20)
        made_by = (m.host << 8) | (45 if zip64 else 20)
        if m.host == HOST_UNIX:
            external = ((0o40755 << 16) | 0x10) if m.directory else (0o100644 << 16)
        else:
            external = 0x10 if m.directory else 0x20
        offset = len(out)

        local_crc, local_c, local_u = (0, 0, 0) if m.descriptor else (crc, len(body), len(m.data))
        local_extra = aes_extra
        if zip64:
            local_extra = _extra(1, struct.pack("<QQ", len(m.data), len(body))) + local_extra
            local_c = local_u = 0xFFFFFFFF
        out += struct.pack("<IHHHHHIIIHH", 0x04034B50, needed, m.flags, method, m.time,
                           ZIP_DOS_DATE, local_crc, local_c, local_u, len(m.name),
                           len(local_extra))
        out += m.name + local_extra + body
        if m.descriptor:
            out += struct.pack("<IIII", 0x08074B50, crc, len(body), len(m.data))

        central_extra = m.extra + aes_extra
        c_size, u_size, c_offset = len(body), len(m.data), offset
        if zip64:
            central_extra = _extra(1, struct.pack("<QQQ", len(m.data), len(body), offset)) + central_extra
            c_size = u_size = c_offset = 0xFFFFFFFF
        central += struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, made_by, needed, m.flags,
                               method, m.time, ZIP_DOS_DATE, crc, c_size, u_size, len(m.name),
                               len(central_extra), 0, 0, 0, external, c_offset)
        central += m.name + central_extra

    index_at = len(out)
    out += central
    count = len(members)
    if zip64:
        record_at = len(out)
        out += struct.pack("<IQHHIIQQQQ", 0x06064B50, 44, 45, 45, 0, 0, count, count,
                           len(central), index_at)
        out += struct.pack("<IIQI", 0x07064B50, 0, record_at, 1)
        out += struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, 0xFFFF, 0xFFFF,
                           0xFFFFFFFF, 0xFFFFFFFF, len(comment))
    else:
        out += struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, count, count, len(central),
                           index_at, len(comment))
    return bytes(out + comment)


def zips() -> None:
    pdf = (OUT / "six-pages.pdf").read_bytes()
    png = (OUT / "tiny.png").read_bytes()
    notes = (OUT / "notes.md").read_bytes()
    plain = (OUT / "plain.txt").read_bytes()
    inner = zip_bytes([Member(b"inside.txt", b"A file inside a zip inside a zip.\n")])
    # Extended timestamp, 2026-01-01T00:00:00Z: the one entry whose time has a zone
    utc = _extra(0x5455, struct.pack("<BI", 1, 1767225600))

    # What people actually have: folders with and without their own entries, a
    # compressed PDF, a stored photo, a file whose sizes trail its data, the
    # clutter macOS adds, zips inside the zip both ways, and the two kinds of file
    # that are listed and cannot be opened. Plus a comment, which moves the end
    # record away from the end of the file.
    archive = [
        Member(b"reports/", method=STORED, directory=True),
        Member(b"reports/six-pages.pdf", pdf, extra=utc),
        Member(b"reports/notes.md", notes, descriptor=True),
        Member(b"photos/tiny.png", png, method=STORED),
        Member(b"photos/.hidden-thumbs", b"not for people"),
        Member(b"__MACOSX/photos/._tiny.png", b"resource fork"),
        Member(b"plain.txt", plain),
        Member(b"nested/stored.zip", inner, method=STORED),
        Member(b"nested/packed.zip", inner),
        Member(b"private/locked.txt", b"The password was gander.\n", zipcrypto=b"gander"),
        Member(b"private/table.dat", bytes(range(40)), method=14),
    ]
    (OUT / "archive.zip").write_bytes(zip_bytes(archive, comment=b"Gander test archive"))
    written(OUT / "archive.zip")

    # Names that would climb out of a folder, start at a root, or use Windows'
    # separator, and one that reorders itself to look like another file.
    odd = [
        Member(b"../escaped.txt", b"one"),
        Member(b"/absolute/path.txt", b"two"),
        Member(b"docs\\readme.txt", b"three", host=HOST_DOS),
        Member(b"unix\\name.txt", b"four"),
        Member(b"a//b/./c.txt", b"five"),
        Member(b"dup.txt", b"first"),
        Member(b"dup.txt", b"second"),
        Member("photo\u202Egpj.apk".encode(), b"six", flags=0x800),
    ]
    (OUT / "odd-names.zip").write_bytes(zip_bytes(odd))
    written(OUT / "odd-names.zip")

    (OUT / "zip64.zip").write_bytes(zip_bytes(
        [Member(b"big/report.pdf", pdf), Member(b"big/notes.txt", plain)], zip64=True))
    written(OUT / "zip64.zip")

    # Every way a file can be under a password that Gander reads, and the one it does not,
    # all with the password gander, beside a file with none. A half past two in the
    # afternoon on the file whose CRC trails it, since that is what its check byte is.
    half_past_two = (14 << 11) | (30 << 5)
    (OUT / "locked.zip").write_bytes(zip_bytes([
        Member(b"open.txt", plain),
        Member(b"zipcrypto.txt", plain, zipcrypto=b"gander"),
        Member(b"zipcrypto-trailing.md", notes, zipcrypto=b"gander", descriptor=True,
               time=half_past_two),
        Member(b"zipcrypto.png", png, method=STORED, zipcrypto=b"gander"),
        Member(b"aes128.txt", plain, aes=(1, 1), password=b"gander"),
        Member(b"aes192.pdf", pdf, aes=(2, 2), password=b"gander"),
        Member(b"aes256.png", png, method=STORED, aes=(3, 2), password=b"gander"),
        Member(b"aes256-deflate64.md", notes, method=DEFLATE64, aes=(3, 2), password=b"gander"),
        Member(b"strong.bin", bytes(range(64)), method=STORED, flags=0x41),
    ]))
    written(OUT / "locked.zip")

    # A password that is not ASCII, as each encryption takes it: WinZip's AES as UTF-8, and
    # the older one as a Russian Windows machine's own code page.
    (OUT / "locked-cyrillic.zip").write_bytes(zip_bytes([
        Member(b"zipcrypto.txt", plain, zipcrypto="пароль".encode("cp866")),
        Member(b"aes.txt", plain, aes=(3, 2), password="пароль".encode("utf-8")),
    ]))
    written(OUT / "locked-cyrillic.zip")

    # Deflate64, which Windows writes for anything over 2 GB. Text to compress, then the
    # three things only Deflate64 has: a match from more than 32 KB back, one from more than
    # 48 KB back, and matches far longer than 258.
    rng = random.Random(64)
    words = [w.encode() for w in (plain + notes).decode().split() if w.isalpha()]
    prose = b" ".join(rng.choice(words) for _ in range(12000))[:70000]
    long_text = (prose + prose[30000:33000] + prose[5000:6000] + b"ab" * 3000 +
                 prose[10000:30000])
    (OUT / "deflate64.zip").write_bytes(zip_bytes([
        Member(b"long.txt", long_text, method=DEFLATE64),
        Member(b"short.txt", plain, method=DEFLATE64),
        Member(b"empty.txt", b"", method=DEFLATE64),
    ]))
    written(OUT / "deflate64.zip")

    # Names as Windows writes them in each code page, flag clear, and as macOS
    # writes them, UTF-8 with the flag clear too.
    text = b"Plain text inside a zip.\n"
    for fixture, encoding, names in [
        ("names-gbk.zip", "gbk", ["季度报告/会议记录.txt", "照片/北京旅行.txt"]),
        ("names-cp866.zip", "cp866", ["Документы/Отчёт за квартал.txt", "Фото/Москва.txt"]),
        ("names-sjis.zip", "shift_jis", ["資料/報告書.txt", "資料/議事録.txt"]),
        ("names-korean.zip", "cp949", ["문서/분기 보고서.pdf", "사진/제주도 여행.jpg"]),
        ("names-mac.zip", "utf-8", ["Отчёт/报告 résumé.txt"]),
    ]:
        members = [Member(n.encode(encoding), text) for n in names]
        (OUT / fixture).write_bytes(zip_bytes(members))
        written(OUT / fixture)


# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# The three word processor formats prose.html reads itself: OpenDocument text,
# Rich Text and Word 97-2003. Each is written here by hand, from the format,
# rather than by a library, so that the bytes are stable and say exactly what
# the tests rely on. What each one carries is the same short document.
#
# These prove the page's plumbing: order, formatting, tables, pictures, notes,
# the error card. They cannot prove the readers against the world, because a
# file written from the same understanding of a format as the reader agrees
# with it by construction; that was done against LibreOffice's rendering of
# real files from the Apache POI and Tika corpora when the readers were built.
# ---------------------------------------------------------------------------

PROSE_TITLE = "Field Survey, Willowmere"
PROSE_BODY = "A short report written only so that a test has something to render."
PROSE_BOLD = "Bold words"
PROSE_AFTER = "The paragraph after it, so ordering is checkable."
PROSE_HINDI = "\u092a\u0941\u0932 \u092c\u0902\u0926 \u0939\u0948\u0964"    # the bridge is closed
PROSE_NOTE = "A footnote, which lands under a rule at the end."


def odt() -> None:
    """letter.odt: a zip of XML, the way LibreOffice writes one, with a picture."""
    png = (OUT / "tiny.png").read_bytes()
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">\n'
        ' <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>\n'
        ' <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="Pictures/tiny.png" manifest:media-type="image/png"/>\n'
        '</manifest:manifest>\n'
    )
    ns = (
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
        'xmlns:xlink="http://www.w3.org/1999/xlink" '
        'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
        'office:version="1.3"'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-styles {ns}>\n'
        ' <office:font-face-decls>\n'
        '  <style:font-face style:name="Liberation Serif" svg:font-family="\'Liberation Serif\'" style:font-family-generic="roman"/>\n'
        ' </office:font-face-decls>\n'
        ' <office:styles>\n'
        '  <style:default-style style:family="paragraph">\n'
        '   <style:text-properties style:font-name="Liberation Serif" fo:font-size="12pt"/>\n'
        '  </style:default-style>\n'
        '  <style:style style:name="Standard" style:family="paragraph"/>\n'
        '  <style:style style:name="Heading_20_1" style:display-name="Heading 1" style:family="paragraph" style:parent-style-name="Standard">\n'
        '   <style:paragraph-properties fo:margin-top="12pt" fo:margin-bottom="6pt"/>\n'
        '   <style:text-properties fo:font-size="18pt" fo:font-weight="bold" fo:color="#1f3864"/>\n'
        '  </style:style>\n'
        ' </office:styles>\n'
        ' <office:automatic-styles>\n'
        '  <style:page-layout style:name="pm1">\n'
        '   <style:page-layout-properties fo:page-width="21cm" fo:page-height="29.7cm" fo:margin-top="2cm" fo:margin-bottom="2cm" fo:margin-left="2cm" fo:margin-right="2cm"/>\n'
        '  </style:page-layout>\n'
        ' </office:automatic-styles>\n'
        ' <office:master-styles>\n'
        '  <style:master-page style:name="Standard" style:page-layout-name="pm1">\n'
        '   <style:header><text:p>HEADER-MARK Willowmere Parish Council</text:p></style:header>\n'
        '  </style:master-page>\n'
        ' </office:master-styles>\n'
        '</office:document-styles>\n'
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-content {ns}>\n'
        ' <office:automatic-styles>\n'
        '  <style:style style:name="T1" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>\n'
        '  <style:style style:name="P1" style:family="paragraph" style:parent-style-name="Standard">\n'
        '   <style:paragraph-properties fo:break-before="page"/>\n'
        '  </style:style>\n'
        '  <style:style style:name="Table1.A1" style:family="table-cell">\n'
        '   <style:table-cell-properties fo:border="0.5pt solid #000000" fo:background-color="#d9e2f3"/>\n'
        '  </style:style>\n'
        '  <text:list-style style:name="L1">\n'
        '   <text:list-level-style-bullet text:level="1" text:bullet-char="\u2022">\n'
        '    <style:list-level-properties><style:list-level-label-alignment text:label-followed-by="listtab" fo:text-indent="-0.635cm" fo:margin-left="1.27cm"/></style:list-level-properties>\n'
        '   </text:list-level-style-bullet>\n'
        '  </text:list-style>\n'
        ' </office:automatic-styles>\n'
        ' <office:body>\n'
        '  <office:text>\n'
        f'   <text:h text:style-name="Heading_20_1" text:outline-level="1">{PROSE_TITLE}</text:h>\n'
        f'   <text:p text:style-name="Standard">{PROSE_BODY}</text:p>\n'
        f'   <text:p text:style-name="Standard"><text:span text:style-name="T1">{PROSE_BOLD}</text:span> in the middle of a sentence.'
        f'<text:note text:id="ftn1" text:note-class="footnote"><text:note-citation>1</text:note-citation>'
        f'<text:note-body><text:p>{PROSE_NOTE}</text:p></text:note-body></text:note></text:p>\n'
        '   <text:list text:style-name="L1">\n'
        '    <text:list-item><text:p>Bridges and culverts</text:p></text:list-item>\n'
        '    <text:list-item><text:p>Retaining walls</text:p></text:list-item>\n'
        '   </text:list>\n'
        '   <table:table table:name="Table1">\n'
        '    <table:table-column table:number-columns-repeated="2"/>\n'
        '    <table:table-row><table:table-cell table:style-name="Table1.A1"><text:p>Item</text:p></table:table-cell>'
        '<table:table-cell table:style-name="Table1.A1"><text:p>Amount</text:p></table:table-cell></table:table-row>\n'
        '    <table:table-row><table:table-cell><text:p>Surveying</text:p></table:table-cell>'
        '<table:table-cell><text:p>4200</text:p></table:table-cell></table:table-row>\n'
        '   </table:table>\n'
        '   <text:p text:style-name="Standard"><draw:frame draw:name="Picture" text:anchor-type="as-char" svg:width="2cm" svg:height="2cm">'
        '<draw:image xlink:href="Pictures/tiny.png" xlink:type="simple"/></draw:frame></text:p>\n'
        f'   <text:p text:style-name="P1">{PROSE_AFTER}</text:p>\n'
        f'   <text:p text:style-name="Standard">Hindi: {PROSE_HINDI}</text:p>\n'
        '  </office:text>\n'
        ' </office:body>\n'
        '</office:document-content>\n'
    )
    members = [
        Member(b"mimetype", b"application/vnd.oasis.opendocument.text", method=STORED),
        Member(b"META-INF/manifest.xml", manifest.encode()),
        Member(b"styles.xml", styles.encode()),
        Member(b"content.xml", content.encode()),
        Member(b"Pictures/tiny.png", png, method=STORED),
    ]
    (OUT / "letter.odt").write_bytes(zip_bytes(members))
    written(OUT / "letter.odt")


def rtf() -> None:
    """memo.rtf: what Word writes, with a code page, Unicode escapes and a picture."""
    png = (OUT / "tiny.png").read_bytes()

    def uni(text: str) -> str:
        # Each character outside ASCII as \uN, with the one-byte fallback \uc1 asks for
        return "".join(c if ord(c) < 128 else f"\\u{ord(c) if ord(c) < 0x8000 else ord(c) - 0x10000}?" for c in text)

    doc = (
        "{\\rtf1\\ansi\\ansicpg1252\\deff0\\deflang1033\\uc1"
        "{\\fonttbl{\\f0\\froman\\fcharset0 Times New Roman;}{\\f1\\fswiss\\fcharset0 Arial;}{\\f2\\fnil\\fcharset2 Wingdings;}}\n"
        "{\\colortbl;\\red31\\green56\\blue100;\\red255\\green255\\blue0;}\n"
        "{\\stylesheet{\\s0 Normal;}{\\s1\\f1\\fs36\\b\\cf1 heading 1;}}\n"
        "{\\*\\generator Gander fixtures}\n"
        "{\\info{\\title Field Survey}}\n"
        "\\paperw11906\\paperh16838\\margl1134\\margr1134\\margt1134\\margb1134\n"
        "{\\header\\pard\\plain\\f0\\fs20 HEADER-MARK Willowmere Parish Council\\par}\n"
        f"\\pard\\plain\\s1\\f1\\fs36\\b\\cf1 {PROSE_TITLE}\\par\n"
        f"\\pard\\plain\\f0\\fs24 {PROSE_BODY}\\par\n"
        f"\\pard\\plain\\f0\\fs24 {{\\b {PROSE_BOLD}}} in the middle of a sentence."
        f"{{\\super\\chftn{{\\footnote\\pard\\plain\\f0\\fs20 {{\\super\\chftn}} {PROSE_NOTE}\\par}}}}\\par\n"
        "{\\pntext\\f2\\'b7\\tab}{\\*\\pn\\pnlvlblt\\pnf2\\pnindent360{\\pntxtb\\'b7}}\\pard\\plain\\fi-360\\li720 Bridges and culverts\\par\n"
        "{\\pntext\\f2\\'b7\\tab}\\pard\\plain\\fi-360\\li720 Retaining walls\\par\n"
        "\\trowd\\trgaph108\\cellx4320\\cellx8640\\clcbpat2"
        "\\pard\\intbl Item\\cell Amount\\cell\\row\n"
        "\\trowd\\trgaph108\\cellx4320\\cellx8640"
        "\\pard\\intbl Surveying\\cell 4200\\cell\\row\n"
        "\\pard\\plain\\f0\\fs24 {\\pict\\pngblip\\picw64\\pich64\\picwgoal1134\\pichgoal1134 " + png.hex() + "}\\par\n"
        "\\page\n"
        f"\\pard\\plain\\f0\\fs24 {PROSE_AFTER}\\par\n"
        f"\\pard\\plain\\f0\\fs24 Hindi: {uni(PROSE_HINDI)}\\par\n"
        "\\pard\\plain\\f0\\fs24 Accents: {\\'e9}t{\\'e9} and Gr{\\'f6}{\\'df}e.\\par\n"
        "}\n"
    )
    (OUT / "memo.rtf").write_bytes(doc.encode("ascii"))
    written(OUT / "memo.rtf")


def doc() -> None:
    """legacy.doc: a Word 97 file, with its text in UTF-16, a bold run and a table.

    The compound file has one FAT sector and one directory sector, and each of
    its two streams is padded to 4,096 bytes so that neither is small enough to
    need the mini stream.
    """
    cell = "\x07"
    parts = [
        PROSE_TITLE + "\r",
        PROSE_BODY + "\r",
        PROSE_BOLD, " in the middle of a sentence.\r",
        "Item" + cell, "Amount" + cell, cell,
        "Surveying" + cell, "4200" + cell, cell,
        PROSE_AFTER + "\r",
        "Hindi: " + PROSE_HINDI + "\r",
    ]
    text = "".join(parts)
    fc_text = 0x400
    fc_end = fc_text + 2 * len(text)
    assert fc_end <= 0x800

    # The paragraphs, and what each one's properties say: nothing, in a table
    # cell, or the end of a row with the row's own definition. The bold words are
    # a run inside the third paragraph, not a paragraph of their own
    para_runs, fc = [], fc_text
    for part in parts:
        end = fc + 2 * len(part)
        if part != PROSE_BOLD:
            start = fc - (2 * len(PROSE_BOLD) if part.startswith(" in the middle") else 0)
            kind = "row" if part == cell else "cell" if part.endswith(cell) else "plain"
            para_runs.append((start, end, kind))
        fc = end

    bold_start = fc_text + 2 * len(parts[0] + parts[1])
    bold_end = bold_start + 2 * len(PROSE_BOLD)

    def u16(v): return struct.pack("<H", v)
    def u32(v): return struct.pack("<I", v)
    def i16(v): return struct.pack("<h", v)

    # CHPX page: three runs, the middle one bold
    chpx = bytes([3, 0x35, 0x08, 1])                       # cb, sprmCFBold, on
    page = bytearray(512)
    for i, f in enumerate([fc_text, bold_start, bold_end, fc_end]):
        page[i * 4:i * 4 + 4] = u32(f)
    page[500:504] = chpx
    page[16:19] = bytes([0, 250, 0])
    page[511] = 3
    chp_page = bytes(page)

    # PAPX page: one run per paragraph, properties at the end of the page
    in_table = bytes([3, 0, 0, 0x16, 0x24, 1])              # cb 3: istd, sprmPFInTable
    tc = u16(0) + u16(4320) + bytes([4, 1, 0, 0]) * 4      # a cell: half-point single borders
    def_table = bytes([2]) + i16(0) + i16(4320) + i16(8640) + tc + tc
    row_end = (u16(0) + bytes([0x16, 0x24, 1, 0x17, 0x24, 1]) +
               bytes([0x08, 0xD6]) + u16(len(def_table) + 1) + def_table)
    assert len(row_end) % 2 == 1
    row_end = bytes([(len(row_end) + 1) // 2]) + row_end
    page = bytearray(512)
    n = len(para_runs)
    for i, (start, _, _) in enumerate(para_runs):
        page[i * 4:i * 4 + 4] = u32(start)
    page[n * 4:n * 4 + 4] = u32(para_runs[-1][1])
    page[400:400 + len(row_end)] = row_end
    page[460:460 + len(in_table)] = in_table
    for i, (_, _, kind) in enumerate(para_runs):
        at = (n + 1) * 4 + i * 13
        page[at] = {"plain": 0, "cell": 230, "row": 200}[kind]
    page[511] = n
    pap_page = bytes(page)

    # The table stream: the piece table, an empty stylesheet, one font, the bin tables
    clx = bytes([2]) + u32(16) + u32(0) + u32(len(text)) + u16(0) + u32(fc_text) + u16(0)
    stsh = u16(18) + u16(0) + u16(10) + u16(0) * 4 + u16(0) * 3
    name = "Times New Roman".encode("utf-16-le") + b"\0\0"
    ffn = bytes([0, 0x12]) + u16(400) + bytes([0, 0]) + bytes(10) + bytes(24) + name
    ffn = bytes([len(ffn) - 1]) + ffn[1:]
    fonts = u16(1) + u16(0) + ffn
    bte_chpx = u32(fc_text) + u32(0x800) + u32(4)
    bte_papx = u32(fc_text) + u32(0x800) + u32(5)
    table = bytearray(4096)
    places = {}
    at = 0
    for key, blob in (("clx", clx), ("stsh", stsh), ("ffn", fonts), ("chpx", bte_chpx), ("papx", bte_papx)):
        table[at:at + len(blob)] = blob
        places[key] = (at, len(blob))
        at += len(blob) + (16 - len(blob) % 16)

    # The FIB, then the text at fc 0x400, then the two formatting pages
    fib = bytearray(0x400)
    fib[0:2] = u16(0xA5EC)
    fib[2:4] = u16(0xC1)
    fib[6:8] = u16(0x0409)
    fib[0x0A:0x0C] = u16(0x0200)                             # fWhichTblStm: 1Table
    fib[0x0C:0x0E] = u16(0xBF)
    fib[0x18:0x1C] = u32(fc_text)
    fib[0x1C:0x20] = u32(fc_end)
    fib[0x20:0x22] = u16(0x0E)
    fib[0x3E:0x40] = u16(0x16)
    fib[0x40:0x44] = u32(fc_end)
    fib[0x4C:0x50] = u32(len(text))
    fib[0x98:0x9A] = u16(0x5D)
    for index, key in ((1, "stsh"), (12, "chpx"), (13, "papx"), (15, "ffn"), (33, "clx")):
        fib[0x9A + index * 8:0x9A + index * 8 + 8] = u32(places[key][0]) + u32(places[key][1])
    word = bytearray(4096)
    word[0:0x400] = fib
    word[0x400:fc_end] = text.encode("utf-16-le")
    word[0x800:0xA00] = chp_page
    word[0xA00:0xC00] = pap_page

    # The compound file around them
    END, FREE, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    header = bytearray(512)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    header[0x18:0x1A] = u16(0x3E)                            # minor version
    header[0x1A:0x1C] = u16(3)
    header[0x1C:0x1E] = u16(0xFFFE)                          # byte order
    header[0x1E:0x20] = u16(9)                               # 512-byte sectors
    header[0x20:0x22] = u16(6)                               # 64-byte mini sectors
    header[0x2C:0x30] = u32(1)                               # one FAT sector
    header[0x30:0x34] = u32(1)                               # directory at sector 1
    header[0x38:0x3C] = u32(4096)
    header[0x3C:0x40] = u32(END)
    header[0x44:0x48] = u32(END)
    header[0x4C:0x50] = u32(0)                               # the FAT is sector 0
    for i in range(1, 109):
        header[0x4C + i * 4:0x50 + i * 4] = u32(FREE)

    fat = [FREE] * 128
    fat[0] = FATSECT
    fat[1] = END
    for s in range(2, 9): fat[s] = s + 1
    fat[9] = END
    for s in range(10, 17): fat[s] = s + 1
    fat[17] = END

    def entry(name: str, kind: int, left, right, child, start, size):
        e = bytearray(128)
        raw = name.encode("utf-16-le") + b"\0\0"
        e[0:len(raw)] = raw
        e[64:66] = u16(len(raw))
        e[66] = kind
        e[67] = 1                                            # black
        e[68:72] = u32(left)
        e[72:76] = u32(right)
        e[76:80] = u32(child)
        e[116:120] = u32(start)
        e[120:124] = u32(size)
        return bytes(e)

    directory = (
        entry("Root Entry", 5, NOSTREAM, NOSTREAM, 1, END, 0) +
        entry("WordDocument", 2, NOSTREAM, 2, NOSTREAM, 2, 4096) +
        entry("1Table", 2, NOSTREAM, NOSTREAM, NOSTREAM, 10, 4096) +
        bytes(128)
    )
    out = bytes(header) + b"".join(u32(v) for v in fat) + directory + bytes(word) + bytes(table)
    assert len(out) == 512 * 19
    (OUT / "legacy.doc").write_bytes(out)
    written(OUT / "legacy.doc")


def prose() -> None:
    odt()
    rtf()
    doc()


# ---------------------------------------------------------------------------
# Workbooks: budget.xlsx's three sheets in the other formats SheetJS reads
# ---------------------------------------------------------------------------

# openpyxl writes none of these, so each is built by hand from its specification, as
# legacy.doc is: the .xls from MS-XLS and MS-CFB, the .xlsb from MS-XLSB, the .ods from
# OpenDocument. The third sheet gains a line outside Latin-1, which each stores its own
# way: BIFF8 in UTF-16 only for a string that needs it, XLSB always, OpenDocument as XML.
WORKBOOK = [
    ("Summary", SHEET_ROWS),
    ("Detail", [("Note", "Value"), ("Second sheet marker", "detail-sheet")]),
    ("Notes", [("Third sheet marker",), ("Grüße aus Zürich, 東京",)]),
]


def cell_value(value):
    """A cell as given in a workbook's rows: a value alone, or (value, fill) for a filled cell"""
    return value[0] if isinstance(value, tuple) else value


def shared_strings(workbook=WORKBOOK):
    """Every string in [workbook] once, in the order first met, and how many cells hold one."""
    strings, total = [], 0
    for _, rows in workbook:
        for value in (cell_value(v) for row in rows for v in row):
            if not isinstance(value, str):
                continue
            total += 1
            if value not in strings:
                strings.append(value)
    return strings, total


def biff(kind: int, body: bytes = b"") -> bytes:
    return struct.pack("<HH", kind, len(body)) + body


def biff_text(text: str, count_size: int) -> bytes:
    """A BIFF8 string, its length in one byte or two: Latin-1 if every character fits."""
    wide = any(ord(c) > 0xFF for c in text)
    return (struct.pack("<B" if count_size == 1 else "<H", len(text)) + bytes([wide])
            + text.encode("utf-16-le" if wide else "latin-1"))


def one_stream_compound_file(name: str, data: bytes) -> bytes:
    """A compound file holding one stream, padded to 4,096 bytes as legacy.doc's streams are."""
    def u16(v): return struct.pack("<H", v)
    def u32(v): return struct.pack("<I", v)

    data += bytes(max(0, 4096 - len(data)))
    size = len(data)
    data += bytes(-len(data) % 512)
    sectors = len(data) // 512
    assert sectors <= 126, "one FAT sector maps 128 sectors"

    END, FREE, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    header = bytearray(512)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    header[0x18:0x1A] = u16(0x3E)
    header[0x1A:0x1C] = u16(3)
    header[0x1C:0x1E] = u16(0xFFFE)
    header[0x1E:0x20] = u16(9)
    header[0x20:0x22] = u16(6)
    header[0x2C:0x30] = u32(1)
    header[0x30:0x34] = u32(1)
    header[0x38:0x3C] = u32(4096)
    header[0x3C:0x40] = u32(END)
    header[0x44:0x48] = u32(END)
    header[0x4C:0x50] = u32(0)
    for i in range(1, 109):
        header[0x4C + i * 4:0x50 + i * 4] = u32(FREE)

    # Sector 0 is the FAT, 1 the directory, and the stream runs on from 2
    fat = [FATSECT, END] + [s + 1 for s in range(2, 1 + sectors)] + [END]
    fat += [FREE] * (128 - len(fat))

    def entry(entry_name: str, kind: int, child, start, length):
        e = bytearray(128)
        raw = entry_name.encode("utf-16-le") + b"\0\0"
        e[0:len(raw)] = raw
        e[64:66] = u16(len(raw))
        e[66] = kind
        e[67] = 1
        e[68:80] = u32(NOSTREAM) + u32(NOSTREAM) + u32(child)
        e[116:124] = u32(start) + u32(length)
        return bytes(e)

    directory = (entry("Root Entry", 5, 1, END, 0) + entry(name, 2, NOSTREAM, 2, size)
                 + bytes(256))
    return bytes(header) + b"".join(u32(v) for v in fat) + directory + data


def biff8(workbook, styles=(), rows_styled=None, columns_styled=None) -> bytes:
    """
    A BIFF8 workbook stream of [workbook]'s sheets, as Excel 97 to 2003 saves one. A cell given
    as (value, k) has the style styles[k]: a fill's pattern and its colour's number in Excel's
    palette, and the colour's number for red bold text or None. One whose value is None is written
    blank, as Excel keeps a blank cell that has a style, and a run of them as one MULBLANK, as
    Excel writes them; one given as ... is not written at all. Each sheet's rows and columns in
    [rows_styled] and [columns_styled] have a style for the cells not written in them, k as above.
    """
    strings, total = shared_strings(workbook)
    rows_styled, columns_styled = rows_styled or {}, columns_styled or {}

    def bof(kind: int) -> bytes:
        return biff(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0, 6))

    def font_record(colour: int = 0x7FFF, weight: int = 400) -> bytes:
        return biff(0x0031, struct.pack("<HHHHHBBBB", 200, 0, colour, weight, 0, 0, 0, 0, 0)
                    + biff_text("Arial", 1))

    def cell_xf(pattern: int = 0, colour: int = 64, red: bool = False) -> bytes:
        # The pattern is the top six bits of the second border word and its colour the low seven
        # of the last, over 65, the window's colour; with a pattern the XF says it sets one. The
        # red bold font is the fifth written, which an XF numbers 5, since no font is number 4.
        used = (0x40 if pattern else 0) | (0x08 if red else 0)
        return biff(0x00E0, struct.pack("<HHHBBBBIIH", 5 if red else 0, 0, 0x0001, 0x20, 0, 0, used,
                                        0, pattern << 26, (65 << 7) | colour))

    fonts = font_record() * 4 + (font_record(10, 700) if any(red for *_, red in styles) else b"")
    # Fifteen style XFs, then the cell XF, number fifteen, that every cell here names but a
    # styled one, which names the XF of its style after it
    style_xf = biff(0x00E0, struct.pack("<HHHBBBBIIH", 0, 0, 0xFFF5, 0x20, 0, 0, 0, 0, 0, 0x20C0))
    head = (bof(0x0005)
            + biff(0x0042, struct.pack("<H", 1200))                        # CODEPAGE: UTF-16
            + biff(0x003D, struct.pack("<HHHHHHHHH", 0, 0, 0x3000, 0x2000, 0x38, 0, 0, 1, 0x258))
            + fonts + style_xf * 15 + cell_xf()
            + b"".join(cell_xf(pattern, colour, bool(red)) for pattern, colour, red in styles)
            + biff(0x0293, struct.pack("<HBB", 0x8000, 0, 0xFF)))          # STYLE: Normal

    def sheet(name, rows, active: bool) -> bytes:
        cells = b""
        for r, row in enumerate(rows):
            blanks = []

            def write_blanks():
                if len(blanks) == 1:
                    return biff(0x0201, struct.pack("<HHH", r, blanks[0][0], blanks[0][1]))      # BLANK
                if blanks:
                    return biff(0x00BE, struct.pack("<HH", r, blanks[0][0])                       # MULBLANK
                                + b"".join(struct.pack("<H", xf) for _, xf in blanks)
                                + struct.pack("<H", blanks[-1][0]))
                return b""

            for c, given in enumerate(row):
                value = cell_value(given)
                xf = 16 + given[1] if isinstance(given, tuple) else 15
                if value is None:
                    blanks.append((c, xf))
                    continue
                cells += write_blanks()
                blanks = []
                if value is ...:
                    continue
                elif isinstance(value, str):
                    cells += biff(0x00FD, struct.pack("<HHHI", r, c, xf, strings.index(value)))
                else:
                    # An RK, which is how Excel stores a whole number that small
                    cells += biff(0x027E, struct.pack("<HHHI", r, c, xf, (value << 2) | 2))
            cells += write_blanks()
        width = max(len(row) for row in rows)
        window = 0x06B6 if active else 0x00B6
        # COLINFO: the columns, their width and their XF; ROW: fGhostDirty says the row has an XF
        columns = b"".join(biff(0x007D, struct.pack("<HHHHHH", c, c, 2560, 16 + k, 0, 0))
                           for c, k in columns_styled.get(name, {}).items())
        row_records = b"".join(biff(0x0208, struct.pack("<HHHHHHHH", r, 0, width, 255, 0, 0, 0x0180, 16 + k))
                               for r, k in rows_styled.get(name, {}).items())
        return (bof(0x0010)
                + columns
                + biff(0x0200, struct.pack("<IIHHH", 0, len(rows), 0, width, 0))
                + row_records
                + cells
                + biff(0x023E, struct.pack("<HHHHHHHI", window, 0, 0, 64, 0, 0, 0, 0))
                + biff(0x000A))

    sheets = [sheet(name, rows, n == 0) for n, (name, rows) in enumerate(workbook)]
    sst = biff(0x00FC, struct.pack("<II", total, len(strings))
               + b"".join(biff_text(s, 2) for s in strings))

    def bound(offsets) -> bytes:
        return b"".join(biff(0x0085, struct.pack("<IBB", at, 0, 0) + biff_text(name, 1))
                        for at, (name, _) in zip(offsets, workbook))

    # Each BOUNDSHEET says where its sheet's BOF is, so the globals are measured first
    at = len(head) + len(bound([0] * len(sheets))) + len(sst) + len(biff(0x000A))
    offsets = []
    for body in sheets:
        offsets.append(at)
        at += len(body)
    return head + bound(offsets) + sst + biff(0x000A) + b"".join(sheets)


def sheet_layout_xls() -> bytes:
    """
    layout.xls's workbook stream: the layout of layout.xlsx that an .xls gives, in BIFF8 built by
    hand from MS-XLS, for xlsx-styles.js's reader of Excel 97-2003's records. A first sheet, a hidden
    second, a third saved as the one showing and frozen at B3, and a fourth with its gridlines off.
    """
    strings = ["The first sheet", "Kept out of sight", "A title", "Frozen row", "In a hidden column", "In a hidden row",
               "Right", "Middle", "Text that wraps onto the lines below it", "Boxed", "Big serif", "Up",
               "No gridlines here"]

    def bof(kind: int) -> bytes:
        return biff(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0, 6))

    def font(height: int, weight: int, name: str) -> bytes:
        return biff(0x0031, struct.pack("<HHHHHBBBB", height, 0, 0x7FFF, weight, 0, 0, 0, 0, 0) + biff_text(name, 1))

    def xf(font_number=0, alc=0, wrap=0, valign=2, trot=0, lines=0, more=0) -> bytes:
        return biff(0x00E0, struct.pack("<HHHBBBBIIH", font_number, 0, 0x0001, alc | wrap << 3 | valign << 4, trot,
                                        0, 0, lines, more, (65 << 7) | 64))

    style_xf = biff(0x00E0, struct.pack("<HHHBBBBIIH", 0, 0, 0xFFF5, 0x20, 0, 0, 0, 0, 0, 0x20C0))
    # Cell XFs from 15: plain, right, centred across and down, wrapped, a thin bottom border in
    # black (palette 8), the fifth font, which an XF numbers 5, and text turned up
    cell_xfs = [xf(), xf(alc=3), xf(alc=2, valign=1), xf(wrap=1), xf(lines=1 << 12, more=8 << 7), xf(font_number=5),
                xf(trot=90)]
    PLAIN, RIGHT, MIDDLE, WRAP, BOXED, BIG, UP = range(15, 22)
    head = (bof(0x0005)
            + biff(0x0042, struct.pack("<H", 1200))
            # WINDOW1: the third sheet is the one showing
            + biff(0x003D, struct.pack("<HHHHHHHHH", 0, 0, 0x3000, 0x2000, 0x38, 2, 0, 1, 0x258))
            + font(200, 400, "Arial") * 4 + font(400, 700, "Times New Roman")
            + style_xf * 15 + b"".join(cell_xfs)
            + biff(0x0293, struct.pack("<HBB", 0x8000, 0, 0xFF)))

    def label(r, c, xf_number, text):
        return biff(0x00FD, struct.pack("<HHHI", r, c, xf_number, strings.index(text)))

    def window(grbit):
        return biff(0x023E, struct.pack("<HHHHHHHI", grbit, 0, 0, 64, 0, 0, 0, 0))

    cover = bof(0x0010) + label(0, 0, PLAIN, "The first sheet") + window(0x00B6) + biff(0x000A)
    hidden = bof(0x0010) + label(0, 0, PLAIN, "Kept out of sight") + window(0x00B6) + biff(0x000A)
    layout = (bof(0x0010)
              # 12.75 points, Arial 10's own, and 8 characters
              + biff(0x0225, struct.pack("<HH", 0, 255))
              + biff(0x0055, struct.pack("<H", 8))
              # COLINFO: A 20 characters, B hidden, C 6
              + biff(0x007D, struct.pack("<HHHHHH", 0, 0, 20 * 256, PLAIN, 0x0002, 0))
              + biff(0x007D, struct.pack("<HHHHHH", 1, 1, 12 * 256, PLAIN, 0x0003, 0))
              + biff(0x007D, struct.pack("<HHHHHH", 2, 2, 6 * 256, PLAIN, 0x0002, 0))
              + biff(0x0200, struct.pack("<IIHHH", 0, 10, 0, 6, 0))
              # ROW: row 2 30 points set by hand, row 4 hidden; rows 10 and 6 nought high, which MS-XLS
              # does not allow, row 6 by hand
              + biff(0x0208, struct.pack("<HHHHHHHH", 1, 0, 1, 600, 0, 0, 0x0140, PLAIN))
              + biff(0x0208, struct.pack("<HHHHHHHH", 3, 0, 1, 255, 0, 0, 0x0120, PLAIN))
              + biff(0x0208, struct.pack("<HHHHHHHH", 9, 0, 5, 0, 0, 0, 0x0100, PLAIN))
              + biff(0x0208, struct.pack("<HHHHHHHH", 5, 0, 1, 0, 0, 0, 0x0140, PLAIN))
              + label(0, 0, PLAIN, "A title") + label(1, 0, PLAIN, "Frozen row")
              + label(2, 1, PLAIN, "In a hidden column") + label(3, 0, PLAIN, "In a hidden row")
              + label(2, 3, RIGHT, "Right") + label(4, 3, MIDDLE, "Middle")
              + label(2, 4, WRAP, "Text that wraps onto the lines below it")
              + label(9, 4, BOXED, "Boxed") + label(8, 0, BIG, "Big serif") + label(7, 0, UP, "Up")
              # An embedded chart: a stream of its own, BOF to EOF, inside the sheet's and before its
              # WINDOW2, as Excel writes one, which the sheet's records go on after
              + bof(0x0020) + biff(0x1002, struct.pack("<IIII", 0, 0, 200 << 16, 100 << 16)) + biff(0x000A)
              # Shown, selected, frozen at B3 without a split
              + window(0x06B6 | 0x0008 | 0x0100)
              + biff(0x0041, struct.pack("<HHHHBB", 1, 2, 2, 1, 0, 0))
              + biff(0x000A))
    plain = (bof(0x0010) + label(0, 0, PLAIN, "No gridlines here")
             # As the others, less fDspGrid
             + window(0x00B6 & ~0x0002) + biff(0x000A))
    sheets = [cover, hidden, layout, plain]
    names = ["Cover", "Hidden", "Layout", "No grid"]
    total = 13
    sst = biff(0x00FC, struct.pack("<II", total, len(strings)) + b"".join(biff_text(t, 2) for t in strings))

    def bound(offsets) -> bytes:
        return b"".join(biff(0x0085, struct.pack("<IBB", at, 1 if name == "Hidden" else 0, 0) + biff_text(name, 1))
                        for n, (at, name) in enumerate(zip(offsets, names)))

    at = len(head) + len(bound([0] * len(sheets))) + len(sst) + len(biff(0x000A))
    offsets = []
    for body in sheets:
        offsets.append(at)
        at += len(body)
    return head + bound(offsets) + sst + biff(0x000A) + b"".join(sheets)


def xls() -> None:
    """budget.xls, and colours.xls: the styles of colours.xlsx that an .xls can give."""
    (OUT / "budget.xls").write_bytes(one_stream_compound_file("Workbook", biff8(WORKBOOK)))
    written(OUT / "budget.xls")
    # Solid yellow, navy, green and orange, numbers 13, 18, 50 and 52 of Excel's palette, and
    # red bold text. B2 and C2 are blank, a run that Excel writes as one record. Row 5 is green and
    # column D orange where no cell is written; D5, where they cross, is written plain.
    styles = [(1, 13, None), (1, 18, None), (0, 64, "red"), (1, 50, None), (1, 52, None)]
    colours = [("Colours", [
        ("Plain",), (("Yellow", 0), (None, 0), (None, 0)), (("Navy", 1),), (("Red bold", 2),),
        (("Whole row", 3), ..., ..., "Crossing"),
    ])]
    stream = biff8(colours, styles, rows_styled={"Colours": {4: 3}}, columns_styled={"Colours": {3: 4}})
    (OUT / "colours.xls").write_bytes(one_stream_compound_file("Workbook", stream))
    written(OUT / "colours.xls")
    (OUT / "layout.xls").write_bytes(one_stream_compound_file("Workbook", sheet_layout_xls()))
    written(OUT / "layout.xls")


def brt(kind: int, body: bytes = b"") -> bytes:
    """A BIFF12 record. Its type and size are each written seven bits a byte, low first."""
    def varint(v: int) -> bytes:
        out = bytearray()
        while True:
            out.append((v & 0x7F) | (0x80 if v > 0x7F else 0))
            v >>= 7
            if not v:
                return bytes(out)
    return varint(kind) + varint(len(body)) + body


def wide_text(text: str) -> bytes:
    """An XLWideString: a count of UTF-16 units, then the units."""
    raw = text.encode("utf-16-le")
    return struct.pack("<I", len(raw) // 2) + raw


OOXML_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def xlsb() -> None:
    """budget.xlsb: Excel's binary workbook, the same parts as an .xlsx with records for XML."""
    strings, total = shared_strings()

    def sheet(rows) -> bytes:
        width = max(len(row) for row in rows)
        out = brt(0x81) + brt(0x94, struct.pack("<IIII", 0, len(rows) - 1, 0, width - 1))
        out += brt(0x91)
        for r, row in enumerate(rows):
            # BrtRowHdr: the row, its style, a height of 15pt, no flags, and one span of cells
            out += brt(0x00, struct.pack("<IIHBBBI", r, 0, 300, 0, 0, 0, 1)
                       + struct.pack("<II", 0, len(row) - 1))
            for c, value in enumerate(row):
                cell = struct.pack("<II", c, 0)
                if isinstance(value, str):
                    out += brt(0x07, cell + struct.pack("<I", strings.index(value)))   # BrtCellIsst
                else:
                    out += brt(0x02, cell + struct.pack("<I", (value << 2) | 2))       # BrtCellRk
        return out + brt(0x92) + brt(0x82)

    book = brt(0x83) + brt(0x8F)
    for n, (name, _) in enumerate(WORKBOOK, start=1):
        book += brt(0x9C, struct.pack("<II", 0, n) + wide_text(f"rId{n}") + wide_text(name))
    book += brt(0x90) + brt(0x84)

    sst = brt(0x9F, struct.pack("<II", total, len(strings)))
    sst += b"".join(brt(0x13, b"\x00" + wide_text(s)) for s in strings)
    sst += brt(0xA0)

    # The least of a style sheet: one font, one cell style and one cell format, which every
    # cell names. Excel always writes one, and SheetJS looks each cell's format up in it.
    def xf(parent: int) -> bytes:
        """A BrtXF: its parent, then font, fill and border 0, and locked, as cells are."""
        return brt(0x2F, struct.pack("<HHHHHBBBBH", parent, 0, 0, 0, 0, 0, 0, 0, 0x10, 0))

    styles = (
        brt(0x116)
        + brt(0x263, struct.pack("<I", 1))
        + brt(0x2B, struct.pack("<HHHHBBBB", 220, 0, 400, 0, 0, 2, 0, 0) + bytes(8)
              + bytes([2]) + wide_text("Calibri"))
        + brt(0x264)
        + brt(0x272, struct.pack("<I", 1)) + xf(0xFFFF) + brt(0x273)
        + brt(0x269, struct.pack("<I", 1)) + xf(0) + brt(0x26A)
        + brt(0x26B, struct.pack("<I", 1))
        + brt(0x30, struct.pack("<IHBB", 0, 1, 0, 0xFF) + wide_text("Normal"))
        + brt(0x26C)
        + brt(0x117)
    )

    main = "application/vnd.ms-excel.sheet.binary.macroEnabled.main"
    types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        f'<Default Extension="bin" ContentType="{main}"/>'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/xl/workbook.bin" ContentType="{main}"/>'
        + "".join(f'<Override PartName="/xl/worksheets/sheet{n}.bin" '
                  'ContentType="application/vnd.ms-excel.worksheet"/>'
                  for n in range(1, len(WORKBOOK) + 1))
        + '<Override PartName="/xl/sharedStrings.bin" ContentType="application/vnd.ms-excel.sharedStrings"/>'
        '<Override PartName="/xl/styles.bin" ContentType="application/vnd.ms-excel.styles"/>'
        '</Types>'
    )
    package_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{OOXML_RELS}/officeDocument" Target="xl/workbook.bin"/>'
        '</Relationships>'
    )
    book_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(f'<Relationship Id="rId{n}" Type="{OOXML_RELS}/worksheet" '
                  f'Target="worksheets/sheet{n}.bin"/>' for n in range(1, len(WORKBOOK) + 1))
        + f'<Relationship Id="rId{len(WORKBOOK) + 1}" Type="{OOXML_RELS}/sharedStrings" '
        'Target="sharedStrings.bin"/>'
        f'<Relationship Id="rId{len(WORKBOOK) + 2}" Type="{OOXML_RELS}/styles" Target="styles.bin"/>'
        '</Relationships>'
    )
    members = [
        Member(b"[Content_Types].xml", types.encode()),
        Member(b"_rels/.rels", package_rels.encode()),
        Member(b"xl/workbook.bin", book),
        Member(b"xl/_rels/workbook.bin.rels", book_rels.encode()),
        *(Member(f"xl/worksheets/sheet{n}.bin".encode(), sheet(rows))
          for n, (_, rows) in enumerate(WORKBOOK, start=1)),
        Member(b"xl/sharedStrings.bin", sst),
        Member(b"xl/styles.bin", styles),
    ]
    (OUT / "budget.xlsb").write_bytes(zip_bytes(members))
    written(OUT / "budget.xlsb")


def ods() -> None:
    """budget.ods: an OpenDocument spreadsheet, a zip of XML as LibreOffice Calc saves one."""
    ns = (
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'office:version="1.3"'
    )

    def cell(value) -> str:
        if isinstance(value, str):
            return f'<table:table-cell office:value-type="string"><text:p>{value}</text:p></table:table-cell>'
        return (f'<table:table-cell office:value-type="float" office:value="{value}">'
                f'<text:p>{value}</text:p></table:table-cell>')

    tables = "".join(
        f'<table:table table:name="{name}">'
        f'<table:table-column table:number-columns-repeated="{max(len(row) for row in rows)}"/>'
        + "".join("<table:table-row>" + "".join(cell(v) for v in row) + "</table:table-row>"
                  for row in rows)
        + "</table:table>"
        for name, rows in WORKBOOK
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-content {ns}><office:body><office:spreadsheet>{tables}'
        '</office:spreadsheet></office:body></office:document-content>\n'
    )
    # LibreOffice always writes a styles.xml, and SheetJS will not open a spreadsheet without one
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-styles {ns}><office:styles/></office:document-styles>\n'
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">\n'
        ' <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>\n'
        ' <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>\n'
        '</manifest:manifest>\n'
    )
    members = [
        Member(b"mimetype", b"application/vnd.oasis.opendocument.spreadsheet", method=STORED),
        Member(b"META-INF/manifest.xml", manifest.encode()),
        Member(b"styles.xml", styles.encode()),
        Member(b"content.xml", content.encode()),
    ]
    (OUT / "budget.ods").write_bytes(zip_bytes(members))
    written(OUT / "budget.ods")


def workbooks() -> None:
    xls()
    xlsb()
    ods()


# ---------------------------------------------------------------------------
# 3D models: one bracket, as the two kinds of STL
# ---------------------------------------------------------------------------

# An L-shaped bracket in millimetres: a plate 40 long and 5 thick with an upright 5 thick
# and 30 tall at one end, 20 deep. Off the origin on purpose, X from 10, so the viewer has
# to find the model's centre rather than assume it. The outline goes anticlockwise seen
# from the front, which is from -Y looking along +Y.
BRACKET_OUTLINE = [(10, 0), (50, 0), (50, 5), (15, 5), (15, 30), (10, 30)]
BRACKET_DEPTH = (-10, 10)

# The outline in four triangles, as corner indices, with no corner on another's edge
BRACKET_CAP = [(0, 1, 2), (0, 2, 3), (0, 3, 5), (3, 4, 5)]


def bracket_triangles():
    """The bracket's twenty triangles, corners anticlockwise seen from outside."""
    front, back = BRACKET_DEPTH
    pts = BRACKET_OUTLINE
    tris = []
    for a, b, c in BRACKET_CAP:
        tris.append([(pts[i][0], front, pts[i][1]) for i in (a, b, c)])
        tris.append([(pts[i][0], back, pts[i][1]) for i in (a, c, b)])
    for i, (ax, az) in enumerate(pts):
        bx, bz = pts[(i + 1) % len(pts)]
        af, ab = (ax, front, az), (ax, back, az)
        bf, bb = (bx, front, bz), (bx, back, bz)
        tris.append([af, bb, bf])
        tris.append([af, ab, bb])
    return tris


def facet_normal(tri):
    (ax, ay, az), (bx, by, bz), (cx, cy, cz) = tri
    ux, uy, uz = bx - ax, by - ay, bz - az
    vx, vy, vz = cx - ax, cy - ay, cz - az
    n = (uy * vz - uz * vy, uz * vx - ux * vz, ux * vy - uy * vx)
    length = sum(c * c for c in n) ** 0.5
    return tuple(c / length for c in n)


def models() -> None:
    tris = bracket_triangles()

    # Binary, with a header that begins "solid" the way SolidWorks and many others write
    # theirs. That is how a text STL begins too, so a reader that trusts the first five
    # bytes reads this one as text and finds nothing in it.
    header = b"solid bracket, binary, from tests/fixtures/make_fixtures.py".ljust(80, b" ")
    body = bytearray(header + struct.pack("<I", len(tris)))
    for tri in tris:
        body += struct.pack("<3f", *facet_normal(tri))
        for corner in tri:
            body += struct.pack("<3f", *corner)
        body += b"\x00\x00"
    (OUT / "bracket.stl").write_bytes(bytes(body))
    written(OUT / "bracket.stl")

    # The same triangles as text, the way SolidWorks writes that: Windows line endings and
    # every number in exponent form
    lines = ["solid bracket"]
    for tri in tris:
        lines.append("  facet normal %e %e %e" % facet_normal(tri))
        lines.append("    outer loop")
        for corner in tri:
            lines.append("      vertex %e %e %e" % corner)
        lines.append("    endloop")
        lines.append("  endfacet")
    lines.append("endsolid bracket")
    (OUT / "bracket-ascii.stl").write_bytes(("\r\n".join(lines) + "\r\n").encode("ascii"))
    written(OUT / "bracket-ascii.stl")


# ---------------------------------------------------------------------------
# Books: EPUB 3, EPUB 2 and one under a lock, for epub.html. Issue #50.
# ---------------------------------------------------------------------------

EPUB_NEEDLE = "Heronsgate"

EPUB_CONTAINER = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
    ' <rootfiles>\n'
    '  <rootfile full-path="{path}" media-type="application/oebps-package+xml"/>\n'
    ' </rootfiles>\n'
    '</container>\n'
)


def xhtml5(title: str, body: str, *, head: str = "", attrs: str = "") -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<!DOCTYPE html>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en">\n'
        f'<head><title>{title}</title>'
        '<link rel="stylesheet" type="text/css" href="../styles/book.css"/>'
        f'{head}</head>\n'
        f'<body{attrs}>\n{body}\n</body>\n</html>\n'
    ).encode("utf-8")


def epub3() -> None:
    """
    harbour.epub: an EPUB 3 in four chapters, made of what a reader has to get right.

    A stylesheet that imports another from its own folder, a picture named from a chapter, a
    stylesheet and an inline style, a font the book scrambles, a link from one chapter to a
    place in another, and a contents list in the nav file with one entry under another.
    Chapter one carries the four ways script gets into a page and a picture on another host.
    Chapter two is long enough to be several pages, and is not sound XML, for an &nbsp; and an
    empty title, so it is the chapter that must be read again as HTML. The needle is once in
    chapter one and twice in chapter three, so a search has to cross chapters to count three.
    """
    buf = io.BytesIO()
    from PIL import Image
    Image.new("RGB", (48, 48), (175, 45, 24)).save(buf, format="PNG")
    mark = buf.getvalue()
    scrambled = random.Random(50).randbytes(600)

    opf = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid">\n'
        ' <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        '  <dc:identifier id="uid">urn:uuid:00000000-0000-4000-8000-000000000050</dc:identifier>\n'
        '  <dc:title>The Harbour at Low Tide</dc:title>\n'
        '  <dc:creator>Imogen Farrow</dc:creator>\n'
        '  <dc:language>en</dc:language>\n'
        f'  <meta property="dcterms:modified">{ISO_EPOCH}</meta>\n'
        ' </metadata>\n'
        ' <manifest>\n'
        '  <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>\n'
        '  <item id="title" href="text/title.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="c1" href="text/chapter-1.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="c2" href="text/chapter-2.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="c3" href="text/chapter-3.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="css" href="styles/book.css" media-type="text/css"/>\n'
        '  <item id="more" href="styles/extra.css" media-type="text/css"/>\n'
        '  <item id="mark" href="images/mark.png" media-type="image/png"/>\n'
        '  <item id="hand" href="fonts/hand.otf" media-type="font/otf"/>\n'
        ' </manifest>\n'
        ' <spine>\n'
        '  <itemref idref="title"/>\n  <itemref idref="c1"/>\n  <itemref idref="c2"/>\n  <itemref idref="c3"/>\n'
        ' </spine>\n'
        '</package>\n'
    )
    nav = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<!DOCTYPE html>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">\n'
        '<head><title>Contents</title></head>\n<body>\n'
        '<nav epub:type="landmarks"><ol><li><a href="text/title.xhtml">Start</a></li></ol></nav>\n'
        '<nav epub:type="toc"><h1>Contents</h1>\n<ol>\n'
        ' <li><a href="text/title.xhtml">Title page</a></li>\n'
        ' <li><a href="text/chapter-1.xhtml">One: The Lamp Room</a>\n'
        '  <ol><li><a href="text/chapter-1.xhtml#ledger">The ledger</a></li></ol></li>\n'
        ' <li><a href="text/chapter-2.xhtml">Two: The Long Watch</a></li>\n'
        ' <li><a href="text/chapter-3.xhtml">Three: The Tide Table</a></li>\n'
        ' <li><span>Notes, in the printed edition only</span></li>\n'
        '</ol></nav>\n</body>\n</html>\n'
    )
    title = xhtml5(
        "The Harbour at Low Tide",
        '<h1 class="book-title">The Harbour at Low Tide</h1>\n'
        '<p class="byline">Imogen Farrow</p>\n'
        '<p><img id="mark" src="../images/mark.png" alt="The printer\'s mark" width="48" height="48"/></p>',
        attrs=' class="front"',
    )
    one = xhtml5(
        "The Lamp Room",
        '<h1>One: The Lamp Room</h1>\n'
        f'<p class="lead">The keeper at {EPUB_NEEDLE} trimmed the lamp at four, as the ledger said he must.</p>\n'
        '<script>window.__xss = "script";</script>\n'
        '<p id="tapped" onclick="window.__xss = \'handler\'">A paragraph with a handler on it.</p>\n'
        '<p><img id="missing" src="nowhere.png" onerror="window.__xss = \'onerror\'" alt="A picture the book does not hold"/></p>\n'
        '<p><img id="remote" src="https://example.invalid/track.png" alt="A picture on another host"/></p>\n'
        '<p><a id="bad-link" href="javascript:window.__xss=\'link\'">A link that is script</a>, '
        '<a id="web-link" href="https://example.invalid/">a web address</a> and '
        '<a id="to-tide" href="chapter-3.xhtml#tide-table">the tide table</a>.</p>\n'
        '<div id="crest" class="crest" style="background-image: url(../images/mark.png); width: 24px; height: 24px"></div>\n'
        '<h2 id="ledger">The ledger</h2>\n'
        '<p class="note">Text after the injected markup, which must survive it.</p>',
    )
    watch = "\n".join(
        f'<p>Hour {n}. The glass stood at {990 + n} and the wind held in the west, so the '
        f'keeper wrote the figure down, wound the clock, and walked the gallery {n} times '
        'before the kettle had boiled.</p>' + ('\n<a id="halfway"/>' if n == 30 else "")
        for n in range(1, 61)
    )
    # Not sound XML on purpose: &nbsp; is not an entity XML knows, and the title is empty
    two = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<!DOCTYPE html>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml" lang="en">\n'
        '<head><title/><link rel="stylesheet" type="text/css" href="../styles/book.css"/></head>\n'
        '<body>\n<h1>Two: The Long&nbsp;Watch</h1>\n' + watch + '\n</body>\n</html>\n'
    ).encode("utf-8")
    three = xhtml5(
        "The Tide Table",
        '<h1>Three: The Tide Table</h1>\n'
        f'<p>Low water at {EPUB_NEEDLE} came at six, and nobody at {EPUB_NEEDLE.upper()} minded.</p>\n'
        '<h2 id="tide-table">The table</h2>\n'
        '<table><tr><th>Day</th><th>Low water</th></tr><tr><td>Monday</td><td>06:10</td></tr></table>\n'
        '<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
        'viewBox="0 0 48 48" width="48" height="48">'
        '<image id="drawn" width="48" height="48" xlink:href="../images/mark.png"/></svg>',
        head='<style>h2 { letter-spacing: 1px; }</style>',
    )
    css = (
        '@import "extra.css";\n'
        ':root { font-style: normal; }\n'
        '@font-face { font-family: "Keeper Hand"; src: url("../fonts/hand.otf"); }\n'
        'body.front { text-align: center; }\n'
        'h1 { color: rgb(120, 40, 20); font-family: "Keeper Hand", serif; }\n'
        'p.lead { text-indent: 0; }\n'
    )
    extra = (
        '.note { font-style: italic; }\n'
        '.byline { background-image: url("../images/mark.png"); background-repeat: no-repeat; padding-left: 56px; }\n'
    )
    encryption = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" '
        'xmlns:enc="http://www.w3.org/2001/04/xmlenc#">\n'
        ' <enc:EncryptedData>\n'
        '  <enc:EncryptionMethod Algorithm="http://www.idpf.org/2008/embedding"/>\n'
        '  <enc:CipherData><enc:CipherReference URI="OEBPS/fonts/hand.otf"/></enc:CipherData>\n'
        ' </enc:EncryptedData>\n'
        '</encryption>\n'
    )
    members = [
        Member(b"mimetype", b"application/epub+zip", method=STORED),
        Member(b"META-INF/container.xml", EPUB_CONTAINER.format(path="OEBPS/package.opf").encode()),
        Member(b"META-INF/encryption.xml", encryption.encode()),
        Member(b"OEBPS/package.opf", opf.encode()),
        Member(b"OEBPS/nav.xhtml", nav.encode()),
        Member(b"OEBPS/text/title.xhtml", title),
        Member(b"OEBPS/text/chapter-1.xhtml", one),
        Member(b"OEBPS/text/chapter-2.xhtml", two),
        Member(b"OEBPS/text/chapter-3.xhtml", three),
        Member(b"OEBPS/styles/book.css", css.encode()),
        Member(b"OEBPS/styles/extra.css", extra.encode()),
        Member(b"OEBPS/images/mark.png", mark, method=STORED),
        Member(b"OEBPS/fonts/hand.otf", scrambled, method=STORED),
    ]
    (OUT / "harbour.epub").write_bytes(zip_bytes(members))
    written(OUT / "harbour.epub")


def epub2() -> None:
    """
    almanac.epub: the older kind of book. Its contents are an NCX file, since it has no nav
    file, its package sits at the top of the zip, and one chapter's name has a space in it,
    which the package writes as %20.
    """
    def page(title: str, body: str) -> bytes:
        return (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.1//EN" "http://www.w3.org/TR/xhtml11/DTD/xhtml11.dtd">\n'
            '<html xmlns="http://www.w3.org/1999/xhtml" xml:lang="en">\n'
            f'<head><title>{title}</title><link rel="stylesheet" type="text/css" href="style.css"/></head>\n'
            f'<body>\n{body}\n</body>\n</html>\n'
        ).encode("utf-8")

    opf = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="uid">\n'
        ' <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        '  <dc:identifier id="uid">urn:uuid:00000000-0000-4000-8000-000000000051</dc:identifier>\n'
        '  <dc:title>A Willowmere Almanac</dc:title>\n'
        '  <dc:language>en</dc:language>\n'
        ' </metadata>\n'
        ' <manifest>\n'
        '  <item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>\n'
        '  <item id="one" href="chapter%20one.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="two" href="two.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="css" href="style.css" media-type="text/css"/>\n'
        ' </manifest>\n'
        ' <spine toc="ncx">\n  <itemref idref="one"/>\n  <itemref idref="two"/>\n </spine>\n'
        '</package>\n'
    )
    ncx = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">\n'
        ' <docTitle><text>A Willowmere Almanac</text></docTitle>\n'
        ' <navMap>\n'
        '  <navPoint id="n1" playOrder="1"><navLabel><text>First Light</text></navLabel>'
        '<content src="chapter%20one.xhtml"/></navPoint>\n'
        '  <navPoint id="n2" playOrder="2"><navLabel><text>Second Watch</text></navLabel>'
        '<content src="two.xhtml"/>\n'
        '   <navPoint id="n3" playOrder="3"><navLabel><text>The watch bell</text></navLabel>'
        '<content src="two.xhtml#bell"/></navPoint>\n'
        '  </navPoint>\n'
        ' </navMap>\n'
        '</ncx>\n'
    )
    members = [
        Member(b"mimetype", b"application/epub+zip", method=STORED),
        Member(b"META-INF/container.xml", EPUB_CONTAINER.format(path="content.opf").encode()),
        Member(b"content.opf", opf.encode()),
        Member(b"toc.ncx", ncx.encode()),
        Member(b"chapter one.xhtml", page("First Light", '<h1>First Light</h1>\n<p>The almanac opens at dawn.</p>')),
        Member(b"two.xhtml", page("Second Watch", '<h1>Second Watch</h1>\n<p>Before the bell.</p>\n'
                                                   '<h2 id="bell">The watch bell</h2>\n<p>It rang twice.</p>')),
        Member(b"style.css", b"h1 { font-variant: small-caps; }\n"),
    ]
    (OUT / "almanac.epub").write_bytes(zip_bytes(members))
    written(OUT / "almanac.epub")


def epub_locked() -> None:
    """
    locked.epub: a book whose one chapter is encrypted, as a shop's DRM leaves it. The list
    in META-INF/encryption.xml is what says so, and the chapter itself is noise.
    """
    opf = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid">\n'
        ' <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        '  <dc:identifier id="uid">urn:uuid:00000000-0000-4000-8000-000000000052</dc:identifier>\n'
        '  <dc:title>A Book Under Lock</dc:title>\n  <dc:language>en</dc:language>\n'
        ' </metadata>\n'
        ' <manifest><item id="c" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>\n'
        ' <spine><itemref idref="c"/></spine>\n'
        '</package>\n'
    )
    encryption = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" '
        'xmlns:enc="http://www.w3.org/2001/04/xmlenc#">\n'
        ' <enc:EncryptedData>\n'
        '  <enc:EncryptionMethod Algorithm="http://www.w3.org/2001/04/xmlenc#aes256-cbc"/>\n'
        '  <enc:CipherData><enc:CipherReference URI="OEBPS/chapter.xhtml"/></enc:CipherData>\n'
        ' </enc:EncryptedData>\n'
        '</encryption>\n'
    )
    members = [
        Member(b"mimetype", b"application/epub+zip", method=STORED),
        Member(b"META-INF/container.xml", EPUB_CONTAINER.format(path="OEBPS/package.opf").encode()),
        Member(b"META-INF/encryption.xml", encryption.encode()),
        Member(b"OEBPS/package.opf", opf.encode()),
        Member(b"OEBPS/chapter.xhtml", random.Random(52).randbytes(900), method=STORED),
    ]
    (OUT / "locked.epub").write_bytes(zip_bytes(members))
    written(OUT / "locked.epub")


def epub_gazetteer() -> None:
    """
    gazetteer.epub: what real books did to the reader that harbour.epub does not. Each part
    is the shape of something met in a book from Project Gutenberg, Standard Ebooks or the
    EPUB 3 samples.

    Its pages are of a fixed size, as a picture book's are, and its first two chapters say
    they are not. The first has the margins a book sets for a desktop window and a heading
    whose id is "title". Then what is wider than any phone: a box given a width in pixels,
    with a margin, a border and a picture in it; a table and a line of code inside a section
    that is not too wide itself; a formula and a run of code set in a line of words; and a
    table in a box that already scrolls sideways. A note is placed far off to the right, as a
    book hides one. Two switches offer markup no reader here knows and a default, written the
    two ways books write them. The second chapter is set in columns that run down the page
    and offers a stylesheet for rows beside the one for columns. The third is the fixed page,
    600 by 800, with a caption placed by measurements from its corner.
    """
    buf = io.BytesIO()
    from PIL import Image
    Image.new("RGB", (1400, 40), (30, 90, 140)).save(buf, format="PNG")
    chart = buf.getvalue()

    def page(title: str, head: str, body: str, attrs: str = "") -> bytes:
        return (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<!DOCTYPE html>\n'
            '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="en">\n'
            f'<head><title>{title}</title>{head}</head>\n'
            f'<body{attrs}>\n{body}\n</body>\n</html>\n'
        ).encode("utf-8")

    opf = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid" '
        'prefix="rendition: http://www.idpf.org/vocab/rendition/#">\n'
        ' <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        '  <dc:identifier id="uid">urn:uuid:00000000-0000-4000-8000-000000000053</dc:identifier>\n'
        '  <dc:title>A Heronsgate Gazetteer</dc:title>\n'
        '  <dc:language>en</dc:language>\n'
        f'  <meta property="dcterms:modified">{ISO_EPOCH}</meta>\n'
        '  <meta property="rendition:layout">pre-paginated</meta>\n'
        ' </metadata>\n'
        ' <manifest>\n'
        '  <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>\n'
        '  <item id="fares" href="fares.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="columns" href="columns.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="plate" href="plate.xhtml" media-type="application/xhtml+xml"/>\n'
        '  <item id="desk" href="desk.css" media-type="text/css"/>\n'
        '  <item id="down" href="down.css" media-type="text/css"/>\n'
        '  <item id="across" href="across.css" media-type="text/css"/>\n'
        '  <item id="chart" href="chart.png" media-type="image/png"/>\n'
        ' </manifest>\n'
        ' <spine>\n'
        '  <itemref idref="fares" properties="rendition:layout-reflowable"/>\n'
        '  <itemref idref="columns" properties="rendition:layout-reflowable"/>\n'
        '  <itemref idref="plate"/>\n'
        ' </spine>\n'
        '</package>\n'
    )
    nav = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<!DOCTYPE html>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">\n'
        '<head><title>Contents</title></head>\n<body>\n'
        '<nav epub:type="toc"><ol>\n'
        ' <li><a href="fares.xhtml#title">Fares</a></li>\n'
        ' <li><a href="columns.xhtml">Columns</a></li>\n'
        ' <li><a href="plate.xhtml">The plate</a></li>\n'
        '</ol></nav>\n</body>\n</html>\n'
    )
    stops = ["Heronsgate", "Willowmere", "Saltmarsh", "Underhill", "Fenwick", "Marrow Cross",
             "Long Ashby", "Tithe Barn", "Coldwater", "Ferry Lane", "Quay Street", "Lantern Row",
             "North Mole", "South Mole"]
    table = (
        '<table id="fares-table"><tr>' + "".join(f"<th>{name}</th>" for name in stops) + "</tr>\n<tr>"
        + "".join(f'<td>{n} shillings and sixpence</td>' for n in range(1, len(stops) + 1)) + "</tr></table>"
    )
    fares = page(
        "Fares",
        '<link rel="stylesheet" type="text/css" href="desk.css"/>',
        # A paragraph long enough to fill a line, so its width is the room the page gave it
        '<p id="first">' + "The ferry leaves the quay on the hour and the half hour. " * 6 + '</p>\n'
        '<h2 id="title">Fares between the stops</h2>\n'
        '<div id="plate" style="width: 1400px; margin-left: 30px; padding: 0 10px; border: 1px solid gray">'
        '<img id="chart" src="chart.png" alt="The fares as a chart"/></div>\n'
        '<section id="tables">\n' + table + '\n'
        '<pre id="log">' + " ".join(f"{n:03d}-crossing" for n in range(60)) + '</pre>\n</section>\n'
        '<p id="sum">The whole run: <code style="white-space: nowrap">'
        + " + ".join(f"fare[{n}]" for n in range(60)) + '</code></p>\n'
        '<p id="eq"><math xmlns="http://www.w3.org/1998/Math/MathML" display="block"><mrow>'
        + "<mo>+</mo>".join(f"<msub><mi>fare</mi><mn>{n}</mn></msub>" for n in range(60)) + '</mrow></math></p>\n'
        '<div id="own" style="overflow-x: auto"><table id="kept"><tr>'
        + "".join(f"<td>{name}</td>" for name in stops * 2) + '</tr></table></div>\n'
        '<span id="far" style="position: absolute; left: 2000px; width: 300px">A note put out of sight.</span>\n'
        # As the EPUB 3 samples write one, in a namespace of its own and no prefix. DOMPurify
        # knows "switch" as part of a drawing, which this is not, and drops all of it.
        '<switch xmlns="http://www.idpf.org/2007/ops" id="formula">'
        '<case required-namespace="http://www.xml-cml.org/schema">'
        '<chem xmlns="http://www.xml-cml.org/schema"><atom>markup no reader here knows</atom></chem>'
        '</case>'
        '<default><p xmlns="http://www.w3.org/1999/xhtml" id="fallback">Salt and water, written out in words.</p></default>'
        '</switch>\n'
        # And with a prefix, where the words of the unknown markup would be shown as well
        '<epub:switch>'
        '<epub:case required-namespace="http://www.xml-cml.org/schema">'
        '<chem xmlns="http://www.xml-cml.org/schema"><atom>more of what no reader here knows</atom></chem>'
        '</epub:case>'
        '<epub:default><p id="fallback-2">A second default, in words.</p></epub:default>'
        '</epub:switch>\n'
        '<p id="after">The last line of the chapter.</p>',
    )
    columns = page(
        "Columns",
        '<link rel="stylesheet" type="text/css" href="down.css" class="vertical" title="Columns"/>'
        '<link rel="alternate stylesheet" type="text/css" href="across.css" class="horizontal" title="Rows"/>'
        '<style>p.aside { writing-mode: vertical-rl; }</style>',
        '<h1>Columns</h1>\n'
        '<div class="main"><p>Set in columns that run down the page, where the book has its way.</p></div>\n'
        '<p class="aside">A note set sideways by the chapter itself.</p>',
    )
    plate = page(
        "The plate",
        '<meta name="viewport" content="width=600, height=800"/>'
        '<style>body { margin: 0; width: 600px; height: 800px; background-color: rgb(20, 40, 60); }\n'
        '#caption { position: absolute; left: 300px; top: 400px; width: 200px; margin: 0; '
        'font-size: 20px; color: rgb(255, 255, 255); }</style>',
        '<p id="caption">Heronsgate from the north mole</p>',
    )
    desk = (
        'body { margin-left: 6em; margin-right: 16em; background-color: rgb(255, 255, 240); }\n'
        '#title { letter-spacing: 3px; }\n'
        'th, td { white-space: nowrap; padding: 4px 12px; }\n'
        'pre { white-space: pre; }\n'
    )
    down = (
        'html { writing-mode: vertical-rl; }\n'
        '.main { margin-left: 12em; }\n'
    )
    across = '.main { letter-spacing: 2px; }\n'
    members = [
        Member(b"mimetype", b"application/epub+zip", method=STORED),
        Member(b"META-INF/container.xml", EPUB_CONTAINER.format(path="EPUB/package.opf").encode()),
        Member(b"EPUB/package.opf", opf.encode()),
        Member(b"EPUB/nav.xhtml", nav.encode()),
        Member(b"EPUB/fares.xhtml", fares),
        Member(b"EPUB/columns.xhtml", columns),
        Member(b"EPUB/plate.xhtml", plate),
        Member(b"EPUB/desk.css", desk.encode()),
        Member(b"EPUB/down.css", down.encode()),
        Member(b"EPUB/across.css", across.encode()),
        Member(b"EPUB/chart.png", chart, method=STORED),
    ]
    (OUT / "gazetteer.epub").write_bytes(zip_bytes(members))
    written(OUT / "gazetteer.epub")


def epubs() -> None:
    epub3()
    epub2()
    epub_locked()
    epub_gazetteer()


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    print(f"Writing fixtures into {OUT}")
    for step in (pdfs, wasm_decoded_images, docx, raised_runs, word_pages, word_columns, word_unrecorded, word_measured,
                 word_odd_even, word_measured_fit, word_measured_notes, word_measured_sections, word_measured_tails,
                 word_measured_tables, word_colours, word_lines,
                 cell_directions, xlsx, sheet_colours, sheet_layout, pptx,
                 without_app_properties, freeforms, straight_lines, wrapping, weights, inherited_bold, line_breaks, unwrapped,
                 symbol_bullets, placed_by_design, unsized, unreadable, charted, undrawn, grounds, relatives, texts, images, audio,
                 photo_formats, zips, prose, workbooks, models, epubs):
        step()
    total = sum(p.stat().st_size for p in OUT.iterdir() if p.is_file())
    count = sum(1 for p in OUT.iterdir() if p.is_file())
    print(f"\n{count} files, {total:,} bytes total")
    if total > 3 * 1024 * 1024:
        print("WARNING: fixtures exceed the 3 MB budget", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
